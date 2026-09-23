"""Auditable GloPro calculations; all arithmetic uses Decimal, never binary floats.

The input schema is discovered from named columns, not physical positions. In
the standard export J/M/P/S mean litres/supplier basis/customer amount/discount.
Unknown schemas and ambiguous business cases must be reviewed by a person.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


class CalculationError(ValueError):
    """The workbook cannot be safely converted into an activation report."""


DEFAULT_RULES = {
    "supplier": "Новое имя",
    "virtual_diesel_multiplier": "0.99",
    "plastic_diesel_multiplier": "1.005",
    "gasoline_multiplier": "1.01",
    "special_customer_multiplier": "1.02",
    "special_clients": ["ПЕРФЕКТ", "АЛЬФА-ЗАПАД"],
    "card_discount_tolerance": "0.20",
}


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFC", str(value or ""))
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip().casefold().replace("ё", "е")


def _header(value: Any) -> str:
    return re.sub(r"[\s.,:;()\-_]+", " ", _norm(value)).strip()


_ALIASES = {
    "date": {"дата", "дата операции", "дата транзакции", "дата и время", "дата и время операции", "дата время"},
    "service": {"услуга", "товар", "вид топлива", "наименование услуги", "наименование товара", "продукт", "топливо"},
    "litres": {"количество", "количество л", "литраж", "объем", "объем л"},
    "supplier_basis": {"стоимость то"},
    "customer_total": {"стоимость"},
    "discount": {"скидка %", "скидка процент", "скидка проценты"},
    "client": {"клиент", "организация", "фирма", "наименование клиента", "наименование организации"},
    "holder": {"держатель", "держатель карты"},
    "card": {"карта", "номер карты"},
    "operation": {"операция", "тип операции"},
}
_ALIASES = {name: {_header(alias) for alias in aliases} for name, aliases in _ALIASES.items()}
_REQUIRED = {"date", "service", "litres", "supplier_basis", "customer_total", "discount"}


def _columns(row: tuple) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, cell in enumerate(row):
        name = _header(cell)
        for field, aliases in _ALIASES.items():
            if name in aliases:
                if field in result:
                    # A duplicated recognized header must never silently select P/M.
                    raise CalculationError(f"Неоднозначный заголовок колонки: {cell!s}")
                result[field] = index
    return result


def _decimal(value: Any, label: str) -> Decimal:
    if value is None or isinstance(value, bool):
        raise CalculationError(f"Отсутствует числовое значение: {label}")
    if isinstance(value, str):
        value = value.strip().replace("\xa0", "").replace("\u202f", "").replace(" ", "").replace(",", ".").replace("−", "-")
        if value.endswith("%"):
            value = value[:-1]
        if value.startswith("(") and value.endswith(")"):
            value = "-" + value[1:-1]
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise CalculationError(f"Неверное числовое значение: {label}") from None
    if not result.is_finite():
        raise CalculationError(f"Неконечное числовое значение: {label}")
    return result


def _date(value: Any, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        raise CalculationError(f"Не распознана дата: {label}") from None


def company_identity(name: str) -> str:
    """Normalize spelling conventions while retaining the legal entity form."""
    name = _norm(name).upper().replace("–", "-").replace("—", "-")
    name = name.translate(str.maketrans({"«": '"', "»": '"', "“": '"', "”": '"', "„": '"'}))
    return re.sub(r'"\s*([^\"]*?)\s*"', lambda match: '"' + match[1].strip() + '"', name)


WEEKLY_CLIENTS = {
    "78749": {"name": "Китай", "kind": "china"},
    "78756": {"name": 'ООО "НК АРТЭЛЬ"', "kind": "nk_artel"},
}


def weekly_kind(client: str, client_id: str | None = None) -> str | None:
    """Exact identities only; a known portal ID takes precedence over its name."""
    if client_id not in (None, ""):
        return WEEKLY_CLIENTS.get(str(client_id), {}).get("kind")
    return next((item["kind"] for item in WEEKLY_CLIENTS.values()
                 if company_identity(client) == company_identity(item["name"])), None)


def _source_clients(rows, header_index, columns):
    clients = set()
    for index, row in enumerate(rows):
        for column, value in enumerate(row):
            if not isinstance(value, str):
                continue
            match = re.fullmatch(r"\s*(?:Клиент|Фирма|Организация)\s*:\s*(.+)\s*", value, re.IGNORECASE)
            if match:
                clients.add(match[1].strip())
            elif index < header_index:
                if re.fullmatch(r'\s*(?:ООО|АО|ПАО|ЗАО|ОАО|ИП)\s+.+', value) or weekly_kind(value):
                    clients.add(value.strip())
                elif _header(value) in _ALIASES["client"]:
                    following = next((str(v).strip() for v in row[column + 1:] if v is not None and str(v).strip()), None)
                    if following:
                        clients.add(following)
        if index > header_index and "client" in columns:
            try:
                _date(row[columns["date"]], "операция")
            except CalculationError:
                continue
            if row[columns["client"]]:
                clients.add(str(row[columns["client"]]).strip())
    return clients


def _source_table(workbook):
    candidates = []
    sheets = [workbook["transactions"]] if "transactions" in workbook.sheetnames else list(workbook)
    for sheet in sheets:
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows[:80]):
            try:
                columns = _columns(row)
            except CalculationError:
                if sum(_header(c) in set.union(*_ALIASES.values()) for c in row if c is not None) >= 4:
                    raise
                continue
            if _REQUIRED - {"discount"} <= columns.keys():
                candidates.append((sheet.title, rows, index, columns))
                break
    if len(candidates) != 1:
        raise CalculationError("Не найден единственный лист с колонками Дата, Услуга, Количество, Стоимость ТО, Стоимость, Скидка %")
    return candidates[0]


def workbook_client(path: str | Path) -> str:
    """Read the source identity before selecting a period for offline imports."""
    try:
        book = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise CalculationError(f"Не удалось прочитать XLSX ({type(exc).__name__})") from exc
    try:
        _, rows, header_index, columns = _source_table(book)
        clients = _source_clients(rows, header_index, columns)
        if len({company_identity(c) for c in clients}) > 1:
            raise CalculationError("В файле указана другая фирма или несколько фирм; проверьте выгрузку")
        return sorted(clients)[0] if clients else ""
    finally:
        book.close()


def _company_key(name: str) -> str:
    """Short name for explicitly configured special-client rules only."""
    name = company_identity(name)
    name = re.sub(r'^(ООО|АО|ПАО|ЗАО|ОАО|ИП)\s+', "", name)
    return name.strip(' \"«»“”„').strip()


def _fuel(value: Any, row_number: int) -> str:
    value = _norm(value).upper().replace("–", "-").replace("—", "-")
    if value in {"ДТ", "ДТ ЭКТО"}:
        return "ДТ ЭКТО"
    value = re.sub(r"^АИ\s*[- ]\s*(\d)", r"АИ-\1", value)
    value = re.sub(r"^ЭКТО\s*[- ]\s*(\d)", r"ЭКТО-\1", value)
    if re.fullmatch(r"АИ-(?:92|95|98|100)(?: ЭКТО| ПРЕМИУМ)?|ЭКТО-(?:92|95|98|100)", value):
        return value
    raise CalculationError(f"Строка {row_number}: неизвестная услуга {value!r}; требуется правило расчёта")


def _summary_row(row: tuple, columns: dict[str, int]) -> bool:
    # Check labels, never numeric equality with a prior total: equal transactions
    # are legitimate and must not be deduplicated.
    if "date" in columns and columns["date"] < len(row):
        try:
            _date(row[columns["date"]], "строка")
            return False
        except CalculationError:
            pass
    for value in row:
        if isinstance(value, str) and re.match(r"^(итого|всего|общий итог)(?:\s|:|$)", _norm(value)):
            return True
    return False


def _money(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _litres(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def _validate_rules(rules: dict | None) -> dict:
    result = {**DEFAULT_RULES, **(rules or {})}
    for key in ("virtual_diesel_multiplier", "plastic_diesel_multiplier", "gasoline_multiplier", "special_customer_multiplier", "card_discount_tolerance"):
        result[key] = _decimal(result[key], key)
        if result[key] <= 0:
            raise CalculationError(f"Правило {key} должно быть положительным")
    if result["card_discount_tolerance"] >= Decimal("0.5"):
        raise CalculationError("Допуск скидки должен быть меньше 0,5 процентного пункта")
    if not isinstance(result["special_clients"], list) or not all(isinstance(s, str) and s.strip() for s in result["special_clients"]):
        raise CalculationError("special_clients должен содержать список названий фирм")
    if result.get("card_type") not in (None, "virtual", "plastic"):
        raise CalculationError("card_type: допустимы virtual или plastic")
    if not isinstance(result["supplier"], str) or not result["supplier"].strip():
        raise CalculationError("Не указано название поставщика")
    for key in ("date_from", "date_to"):
        if result.get(key):
            result[key] = _date(result[key], key)
    if result.get("date_from") and result.get("date_to") and result["date_from"] > result["date_to"]:
        raise CalculationError("Начало периода позже окончания")
    return result


def calculate_workbook(path: str | Path, rules: dict | None = None) -> dict:
    """Calculate one company's XLSX report.

    `rules` accepts DEFAULT_RULES overrides, `client`, optional `card_type`, and
    inclusive `date_from`/`date_to`. A period mismatch rejects the workbook rather
    than silently dropping transactions. Decimal quantities are serialized as
    strings. `status` is `ready`, `no_data`, or `fully_reversed`.
    """
    settings = _validate_rules(rules)
    try:
        # Some genuine GloPro exports advertise dimension A1 although hundreds
        # of XML rows exist. Normal mode reads actual cells instead of truncating.
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise CalculationError(f"Не удалось прочитать XLSX ({type(exc).__name__})") from exc
    try:
        sheet_name, rows, header_index, columns = _source_table(workbook)
        client = str(settings.get("client") or "").strip()
        metadata_clients = _source_clients(rows, header_index, columns)
        client_keys = {company_identity(value) for value in metadata_clients}
        if len(client_keys) > 1 or (client and client_keys and company_identity(client) not in client_keys):
            raise CalculationError("В файле указана другая фирма или несколько фирм; проверьте выгрузку")
        if not client and metadata_clients:
            client = sorted(metadata_clients)[0]
        kind = weekly_kind(client, settings.get("client_id"))
        if kind:
            from .weekly import calculate_weekly
            if sheet_name != "transactions":
                raise CalculationError("Для недельного расчёта нужен лист transactions")
            return calculate_weekly(rows, header_index, columns, client, kind, settings)
        if not _REQUIRED <= columns.keys():
            raise CalculationError("Не найден единственный лист с колонками Дата, Услуга, Количество, Стоимость ТО, Стоимость, Скидка %")
        groups: dict[str, dict] = {}
        transaction_dates = []
        discount_types = set()
        skipped_summaries = []
        skipped_headers = []
        transaction_rows = []
        skipped_metadata = []
        summary_section = False
        for index, row in enumerate(rows[header_index + 1:], header_index + 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if _summary_row(row, columns):
                skipped_summaries.append(index)
                if any(re.match(r"^(итого по отчету|общий итог по услугам)", _norm(value)) for value in row if isinstance(value, str)):
                    summary_section = True
                continue
            if _REQUIRED <= _columns(row).keys():
                skipped_headers.append(index)
                continue
            def cell(field: str):
                return row[columns[field]] if columns[field] < len(row) else None
            if summary_section:
                try:
                    _date(cell("date"), f"строка {index}")
                except CalculationError:
                    skipped_summaries.append(index)
                    continue
                raise CalculationError(f"Строка {index}: транзакция после общего итога; неясная структура отчёта")
            row_metadata = []
            for value in row:
                if isinstance(value, str) and re.match(r"^\s*(Клиент|Фирма|Организация|Карта|Держатель|Услуга)\s*:", value, re.IGNORECASE):
                    row_metadata.append(value)
                    match = re.match(r"^\s*(?:Клиент|Фирма|Организация)\s*:\s*(.+)", value, re.IGNORECASE)
                    if match:
                        metadata_clients.add(match[1].strip())
            def metadata_or_empty(value):
                return value in (None, "") or isinstance(value, str) and re.match(r"^\s*(Клиент|Фирма|Организация|Карта|Держатель|Услуга)\s*:", value, re.IGNORECASE)
            if row_metadata and all(cell(key) in (None, "") for key in ("litres", "supplier_basis", "customer_total")) and all(metadata_or_empty(cell(key)) for key in ("date", "service")):
                skipped_metadata.append(index)
                continue
            # Explicit empty-table notices are distinct from arbitrary malformed data.
            if any(_norm(value) in {"нет данных", "нет операций", "операции отсутствуют", "данные отсутствуют"} for value in row):
                if any(cell(key) not in (None, "") for key in ("litres", "supplier_basis", "customer_total")):
                    raise CalculationError(f"Строка {index}: сообщение об отсутствии данных содержит суммы")
                continue
            fuel = _fuel(cell("service"), index)
            tx_date = _date(cell("date"), f"строка {index}")
            if settings.get("date_from") and tx_date < settings["date_from"] or settings.get("date_to") and tx_date > settings["date_to"]:
                raise CalculationError(f"Строка {index}: операция {tx_date.isoformat()} вне запрошенного периода")
            amounts = {key: _decimal(cell(key), f"{key}, строка {index}") for key in ("litres", "supplier_basis", "customer_total")}
            if amounts["litres"] == 0 and any(amounts[key] != 0 for key in ("supplier_basis", "customer_total")):
                raise CalculationError(f"Строка {index}: нулевой литраж при ненулевой сумме")
            if any(amounts["litres"] * amounts[key] < 0 for key in ("supplier_basis", "customer_total")):
                raise CalculationError(f"Строка {index}: знаки литража и сумм не совпадают")
            if "client" in columns and cell("client"):
                metadata_clients.add(str(cell("client")).strip())
            group = groups.setdefault(fuel, {"litres": Decimal(0), "supplier_basis": Decimal(0), "customer_total": Decimal(0), "rows": []})
            for key, amount in amounts.items():
                group[key] += amount
            group["rows"].append(index)
            transaction_rows.append(index)
            transaction_dates.append(tx_date)
            if fuel == "ДТ ЭКТО" and amounts["litres"] != 0:
                discount = _decimal(cell("discount"), f"скидка, строка {index}")
                # XLSX percentages may be stored as fractions with a '%' number
                # format. Numeric -0.01 alone is not guessed to mean -1%.
                number_format = workbook[sheet_name].cell(index, columns["discount"] + 1).number_format
                if "%" in number_format and not isinstance(cell("discount"), str):
                    discount *= 100
                if abs(discount + 1) <= settings["card_discount_tolerance"]:
                    discount_types.add("virtual")
                elif abs(discount + 2) <= settings["card_discount_tolerance"]:
                    discount_types.add("plastic")
                else:
                    discount_types.add("unknown")
        client_keys = {company_identity(value) for value in metadata_clients}
        if len(client_keys) > 1 or (client and client_keys and company_identity(client) not in client_keys):
            raise CalculationError("В файле указана другая фирма или несколько фирм; проверьте выгрузку")
        if not client and metadata_clients:
            client = sorted(metadata_clients)[0]
        if not client:
            raise CalculationError("Не удалось определить фирму: задайте rules.client или добавьте колонку Клиент")
        is_special = _company_key(client) in {_company_key(name) for name in settings["special_clients"]}
        if len(discount_types - {"unknown"}) > 1:
            raise CalculationError("В дизельных операциях смешаны виртуальные и пластиковые карты; требуется раздельный расчёт")
        if is_special and any(fuel != "ДТ ЭКТО" for fuel in groups):
            raise CalculationError(f"{client}: исключение для бензина не определено; уточните правило клиента")
        if is_special:
            card_type = "virtual"
        elif "ДТ ЭКТО" in groups and groups["ДТ ЭКТО"]["litres"] != 0:
            if not discount_types or "unknown" in discount_types:
                raise CalculationError("Тип дизельной карты не определён по скидке; требуется проверка")
            card_type = next(iter(discount_types))
            if settings.get("card_type") and settings["card_type"] != card_type:
                raise CalculationError("Заданный тип карты противоречит скидке в выгрузке")
        else:
            card_type = settings.get("card_type")
        fuels = []
        raw_fuels = []
        reversed_fuels = []
        totals = {key: Decimal(0) for key in ("litres", "supplier_basis", "customer_total")}
        with localcontext() as context:
            context.prec = 32
            for fuel, group in sorted(groups.items(), key=lambda item: (item[0] != "ДТ ЭКТО", item[0])):
                if group["litres"] == 0:
                    if group["supplier_basis"] != 0 or group["customer_total"] != 0:
                        raise CalculationError(f"{fuel}: итоговый литраж равен нулю при ненулевой сумме")
                    reversed_fuels.append(fuel)
                    raw_fuels.append({"fuel": fuel, **{key: str(group[key]) for key in ("litres", "supplier_basis", "customer_total")}})
                    continue
                multiplier = settings["gasoline_multiplier"]
                if fuel == "ДТ ЭКТО":
                    multiplier = settings[f"{card_type}_diesel_multiplier"]
                    if is_special:
                        group["customer_total"] = group["supplier_basis"] * settings["special_customer_multiplier"]
                price = group["supplier_basis"] / group["litres"] * multiplier
                raw_fuels.append({"fuel": fuel, **{key: str(group[key]) for key in ("litres", "supplier_basis", "customer_total")}})
                fuels.append({"fuel": fuel, "litres": _litres(group["litres"]), "supplier_basis": _money(group["supplier_basis"]), "customer_total": _money(group["customer_total"]), "supplier_price": _money(price), "rows": group["rows"]})
                for key in totals:
                    totals[key] += group[key]
        return {
            "client": client,
            "supplier": settings["supplier"],
            "card_type": card_type,
            "status": "ready" if fuels else "fully_reversed" if transaction_rows else "no_data",
            "fuels": fuels,
            "totals": {key: _litres(value) if key == "litres" else _money(value) for key, value in totals.items()},
            "audit": {"sheet": sheet_name, "header_row": header_index + 1, "columns": {key: index + 1 for key, index in columns.items()}, "transaction_rows": transaction_rows, "skipped_summary_rows": skipped_summaries, "skipped_header_rows": skipped_headers, "skipped_metadata_rows": skipped_metadata, "fully_reversed_fuels": reversed_fuels, "date_from": min(transaction_dates).isoformat() if transaction_dates else None, "date_to": max(transaction_dates).isoformat() if transaction_dates else None, "special_client": is_special, "discount_types": sorted(discount_types), "raw_fuels": raw_fuels, "raw_totals": {key: str(value) for key, value in totals.items()}},
        }
    finally:
        workbook.close()


def merge_reports(reports: Iterable[dict], rules: dict | None = None) -> dict:
    """Merge verified, distinct contracts of one company before rounding.

    The caller must verify unique source/contract identities: two identical
    transactions across independent contracts cannot be distinguished here.
    Supplier unit prices are recomputed from raw totals, never averaged.
    """
    reports = list(reports)
    if not reports:
        raise CalculationError("Нет отчётов для объединения")
    if any(report.get("weekly_kind") for report in reports):
        from .weekly import merge_weekly_reports
        return merge_weekly_reports(reports)
    settings = _validate_rules(rules)
    clients = {company_identity(report["client"]) for report in reports}
    suppliers = {_norm(report["supplier"]) for report in reports}
    if len(clients) != 1 or len(suppliers) != 1:
        raise CalculationError("Нельзя объединять разные фирмы или разных поставщиков")
    client, supplier = reports[0]["client"], reports[0]["supplier"]
    if settings.get("client") and company_identity(settings["client"]) not in clients:
        raise CalculationError("Заданная фирма не совпадает с объединяемыми договорами")
    if rules and "supplier" in rules and _norm(rules["supplier"]) not in suppliers:
        raise CalculationError("Заданный поставщик не совпадает с объединяемыми договорами")
    active = [report for report in reports if report.get("status") == "ready" and report.get("fuels")]
    card_types = {report.get("card_type") for report in active} - {None}
    if card_types - {"virtual", "plastic"}:
        raise CalculationError("Неизвестный тип карты в договорах")
    if len(card_types) > 1:
        raise CalculationError("В договорах смешаны виртуальные и пластиковые карты; требуется проверка")
    is_special = _company_key(client) in {_company_key(name) for name in settings["special_clients"]}
    card_type = "virtual" if is_special else next(iter(card_types), settings.get("card_type"))
    groups = {}
    dates_from, dates_to = [], []
    for source_index, report in enumerate(reports):
        if report.get("status") not in ("ready", "no_data", "fully_reversed"):
            raise CalculationError("Объединять можно только проверенные результаты расчёта")
        audit = report.get("audit", {})
        if audit.get("date_from"):
            dates_from.append(_date(audit["date_from"], "начало договора"))
        if audit.get("date_to"):
            dates_to.append(_date(audit["date_to"], "конец договора"))
        seen_fuels = set()
        source_fuels = audit.get("raw_fuels", report["fuels"])
        if report.get("status") == "no_data" and source_fuels:
            raise CalculationError("Пустой отчёт содержит топливные операции")
        for item in source_fuels:
            fuel = _fuel(item["fuel"], source_index + 1)
            if fuel in seen_fuels:
                raise CalculationError("Повтор топлива внутри договора: неоднозначная группировка")
            seen_fuels.add(fuel)
            group = groups.setdefault(fuel, {"litres": Decimal(0), "supplier_basis": Decimal(0), "customer_total": Decimal(0), "rows": []})
            for key in ("litres", "supplier_basis", "customer_total"):
                group[key] += _decimal(item[key], f"{key}, договор {source_index + 1}")
            group["rows"].append({"report": source_index, "fuel": fuel})
    if is_special and any(fuel != "ДТ ЭКТО" for fuel in groups):
        raise CalculationError(f"{client}: исключение для бензина не определено; уточните правило клиента")
    if settings.get("card_type") and card_type and settings["card_type"] != card_type:
        raise CalculationError("Заданный тип карты противоречит типу карт договоров")
    totals = {key: Decimal(0) for key in ("litres", "supplier_basis", "customer_total")}
    fuels, raw_fuels, reversed_fuels = [], [], []
    with localcontext() as context:
        context.prec = 32
        for fuel, group in sorted(groups.items(), key=lambda item: (item[0] != "ДТ ЭКТО", item[0])):
            if group["litres"] == 0:
                if group["supplier_basis"] != 0 or group["customer_total"] != 0:
                    raise CalculationError(f"{fuel}: итоговый литраж равен нулю при ненулевой сумме")
                reversed_fuels.append(fuel)
                raw_fuels.append({"fuel": fuel, **{key: str(group[key]) for key in totals}})
                continue
            multiplier = settings["gasoline_multiplier"]
            if fuel == "ДТ ЭКТО":
                if not card_type:
                    raise CalculationError("Тип дизельных карт договоров не определён")
                multiplier = settings[f"{card_type}_diesel_multiplier"]
                if is_special:
                    group["customer_total"] = group["supplier_basis"] * settings["special_customer_multiplier"]
            price = group["supplier_basis"] / group["litres"] * multiplier
            raw_fuels.append({"fuel": fuel, **{key: str(group[key]) for key in totals}})
            fuels.append({"fuel": fuel, "litres": _litres(group["litres"]), "supplier_basis": _money(group["supplier_basis"]), "customer_total": _money(group["customer_total"]), "supplier_price": _money(price), "rows": group["rows"]})
            for key in totals:
                totals[key] += group[key]
    return {
        "client": client, "supplier": supplier, "card_type": card_type,
        "status": "ready" if fuels else "fully_reversed" if groups or any(r["status"] == "fully_reversed" for r in reports) else "no_data",
        "fuels": fuels,
        "totals": {key: _litres(value) if key == "litres" else _money(value) for key, value in totals.items()},
        "audit": {"merged": True, "source_reports": reports, "contract_ids": [report.get("contract_id") or report.get("audit", {}).get("contract_id") for report in reports], "date_from": min(dates_from).isoformat() if dates_from else None, "date_to": max(dates_to).isoformat() if dates_to else None, "fully_reversed_fuels": reversed_fuels, "special_client": is_special, "raw_fuels": raw_fuels, "raw_totals": {key: str(value) for key, value in totals.items()}},
    }


def _ru(value: str, places: int) -> str:
    return f"{Decimal(value):,.{places}f}".replace(",", " ").replace(".", ",")


def render_report(reports: Iterable[dict], activation_date: date | str) -> str:
    """Render only the concise customer-facing format from the vault instruction."""
    report_date = _date(activation_date, "дата активации")
    blocks = [report_date.strftime("%d.%m.%Y")]
    from .ordering import alphabet_key
    for report in sorted(reports, key=lambda item: alphabet_key(item["client"])):
        client = str(report["client"]).replace("\n", " ").replace("\r", " ")
        supplier = str(report.get("supplier", DEFAULT_RULES["supplier"])).replace("\n", " ").replace("\r", " ")
        lines = [f"### {client}", f'Поставщик: "{supplier}"  ']
        if report.get("weekly_kind"):
            from .weekly import render_weekly
            blocks.append(render_weekly(report))
            continue
        fuels = report["fuels"]
        if not fuels:
            lines.append("Операции полностью возвращены." if report.get("status") == "fully_reversed" else "Нет операций за период.")
        elif len(fuels) == 1 and fuels[0]["fuel"] == "ДТ ЭКТО":
            fuel = fuels[0]
            lines.extend([f'{_ru(fuel["supplier_price"], 2)} рублей  ', f'{_ru(fuel["litres"], 3)} л | {_ru(fuel["customer_total"], 2)} рублей'])
        else:
            lines.extend(f'- {fuel["fuel"]} — {_ru(fuel["litres"], 3)} л | {_ru(fuel["customer_total"], 2)} руб | {_ru(fuel["supplier_price"], 2)} руб/л' for fuel in fuels)
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks) + "\n"

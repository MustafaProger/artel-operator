"""Strict reconciliation of the Yandex Go corporate fuel XLSX export."""
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import hashlib
import re

from openpyxl import load_workbook, Workbook
from .ordering import alphabet_key


def period_for(run_date):
    if run_date.weekday() != 1:
        raise ValueError("Яндекс актируется только во вторник за предыдущие вторник–понедельник")
    return run_date - timedelta(days=7), run_date - timedelta(days=1)


def text(value):
    return " ".join(str(value or "").split())


def number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("В числовой ячейке Яндекса нет значения")
    try:
        result = Decimal(str(value).replace(" ", "").replace("\xa0", "").replace(",", "."))
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise ValueError("Некорректное числовое значение в отчёте Яндекса") from None


def phone_hash(value):
    return hashlib.sha256(re.sub(r"\D", "", str(value)).encode()).hexdigest()


def display_name(source_name, user_id, phone, aliases):
    matches = [e for e in aliases if (user_id and e.get("user_id") == user_id)
               or (not user_id and e.get("phone_sha256") == phone)]
    if len(matches) > 1:
        raise ValueError("Неоднозначная привязка сотрудника Яндекса")
    return matches[0]["name"] if matches else source_name


def safe_label(value):
    # Keep source spelling in captions, but prevent path traversal/Markdown headings.
    cleaned = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", value).strip(" .")
    return cleaned.encode("utf-8")[:140].decode("utf-8", errors="ignore") or "Сотрудник"


def label_reports(reports):
    for report in reports:
        base = safe_label(report["name"])
        duplicates = [r for r in reports if safe_label(r["name"]).casefold() == base.casefold()]
        report["label"] = base + (f" [{report['employee_id']}]" if len(duplicates) > 1 else "")
    return sorted(reports, key=lambda r: (alphabet_key(r["name"]), r["employee_id"]))


def read_report(path, start, end, company, employees=None, *, aliases=(), expected_orders=None):
    with Path(path).open("rb") as stream:
        book = load_workbook(stream, data_only=False, read_only=True)
        try:
            if "Отчёт" not in book.sheetnames:
                raise ValueError("В файле Яндекса отсутствует лист «Отчёт»")
            sheet = book["Отчёт"]
            if sheet.max_row > 100000 or sheet.max_column > 200:
                raise ValueError("Размер отчёта Яндекса превышает допустимый")
            rows = list(sheet.values)
        finally:
            book.close()
    if len(rows) < 8:
        raise ValueError("Отчёт Яндекса неполный")
    metadata = {text(row[0]): row[3] for row in rows[:6] if len(row) > 3 and row[0]}
    if text(metadata.get("Компания")) != text(company):
        raise ValueError("В отчёте Яндекса другая организация")
    expected_period = f"{start:%d.%m.%Y} - {end:%d.%m.%Y}"
    if text(metadata.get("Период")) != expected_period:
        raise ValueError("Период в Excel Яндекса не совпадает с выбранной неделей")
    if text(metadata.get("Данные указаны по часовому поясу")) != "UTC+3":
        raise ValueError("В отчёте Яндекса должен быть часовой пояс UTC+3")
    headers = [text(c) for c in rows[6]]
    required = ["Дата заказа", "Имя пользователя", "Телефон", "Идентификатор заказа", "Топливо", "Залито", "Статус", "Стоимость"]
    if any(headers.count(name) != 1 for name in required):
        raise ValueError("Выгрузка Яндекса должна содержать дату, сотрудника, телефон, ID заказа, топливо, «Залито», статус и «Стоимость»")
    columns = {name: headers.index(name) for name in required}
    cost_column = columns["Стоимость"]
    operations, ids, found_employees, footer = [], set(), {}, None
    checked_orders = []
    expected = {o["id"]: o for o in expected_orders} if expected_orders is not None else None
    for row in rows[7:]:
        if not any(value is not None for value in row):
            continue
        if text(row[cost_column - 1]) == "Итого":
            if footer is not None or any(value is not None for value in row[:cost_column - 1]):
                raise ValueError("Некорректная итоговая строка Яндекса")
            footer = number(row[cost_column])
            continue
        if footer is not None:
            raise ValueError("После итоговой строки Яндекса найдены операции")
        values = {name: row[index] for name, index in columns.items()}
        name = text(values["Имя пользователя"])
        raw_phone = re.sub(r"\D", "", str(values["Телефон"] or ""))
        if not name or not re.fullmatch(r"\d{10,15}", raw_phone):
            raise ValueError("В Excel Яндекса нет имени или полного телефона сотрудника")
        identity = phone_hash(raw_phone)
        order_id = text(values["Идентификатор заказа"])
        if not re.fullmatch(r"[0-9a-f]{32}", order_id) or order_id in ids:
            raise ValueError("Отсутствует или повторяется ID заказа Яндекса")
        ids.add(order_id)
        if expected is not None:
            order = expected.get(order_id)
            if not order or order["source_name"] != name or order["phone_sha256"] != identity:
                raise ValueError("Сотрудник или ID заказа Excel не совпал с кабинетом Яндекса")
            employee = dict(user_id=order["user_id"], source_name=name, phone_sha256=identity)
            employee["name"] = display_name(name, employee["user_id"], identity, aliases)
        elif employees is not None:
            matches = [e for e in employees if e["source_name"] == name and e["phone_sha256"] == identity]
            if len(matches) != 1:
                raise ValueError("В отчёте Яндекса неизвестный сотрудник или другой одноимённый сотрудник")
            employee = matches[0]
        else:
            # Standard Yandex XLSX has no internal user_id. Do not invent one:
            # offline imports use the full phone fingerprint, explicitly audited.
            matches = [e for e in aliases if e.get("phone_sha256") == identity]
            if len(matches) > 1:
                raise ValueError("Один телефон связан с несколькими ID сотрудников; нужен отчёт с проверкой кабинета")
            user_id = matches[0]["user_id"] if matches else None
            employee = dict(user_id=user_id, source_name=name, phone_sha256=identity,
                            name=display_name(name, user_id, identity, aliases))
        employee_id = employee.get("user_id") or "phone-" + identity
        previous = found_employees.get(employee_id)
        if expected is None and previous and previous["source_name"] != name:
            raise ValueError("В XLSX разные имена на одном телефоне; без user_id нельзя подтвердить одного сотрудника")
        found_employees[employee_id] = employee
        if len(found_employees) != 1:
            raise ValueError("Загружайте отдельный отчёт Яндекса для каждого сотрудника")
        value = values["Дата заказа"]
        try:
            day = value.date() if isinstance(value, datetime) else datetime.strptime(str(value), "%d.%m.%Y").date()
        except (ValueError, TypeError):
            raise ValueError("Не удалось проверить дату операции Яндекса") from None
        if not start <= day <= end:
            raise ValueError("В отчёте Яндекса есть операция вне выбранной недели")
        litres, amount = number(values["Залито"]), number(values["Стоимость"])
        status = text(values["Статус"]).casefold().replace("ё", "е")
        checked_orders.append({"id": order_id, "source_name": name, "phone_sha256": identity,
                               "litres": str(litres), "amount": str(amount), "date": day.isoformat()})
        if expected is not None:
            order = expected[order_id]
            if (amount != Decimal(order["amount"]) or litres != Decimal(order["litres"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                    or day.isoformat() != order["date"]):
                raise ValueError("Строка Excel не совпала с суммой, литрами или датой заказа в кабинете")
        if status in {"отменен", "отменен пользователем"} and litres == amount == 0:
            if expected is not None and expected[order_id]["status"] != "Cancelled":
                raise ValueError("Статус Excel не совпал с кабинетом")
            continue
        if expected is not None and expected[order_id]["status"] != "Completed":
            raise ValueError("Статус Excel не совпал с кабинетом")
        if status != "завершен":
            raise ValueError("Неподдержанный статус операции Яндекса: нужна проверка возврата или незавершённой заправки")
        if not text(values["Топливо"]) or litres == 0 and amount != 0 or litres < 0 or amount < 0:
            raise ValueError("Литры и сумма операции Яндекса не согласуются")
        operations.append({"id": order_id, "date": day.isoformat(), "fuel": text(values["Топливо"]),
                           "litres": str(litres), "amount": str(amount)})
    declared_count = number(metadata.get("Количество заказов в отчётном периоде"))
    if declared_count != len(ids) or footer is None:
        raise ValueError("Количество заказов или итоговая строка Яндекса не совпадают с выгрузкой")
    total = sum((Decimal(o["amount"]) for o in operations), Decimal(0))
    if total != footer or total != number(metadata.get("Общая стоимость с НДС")):
        raise ValueError("Сумма операций Яндекса не совпадает с итогами файла")
    if not found_employees:
        raise ValueError("В пустом файле Яндекса невозможно подтвердить сотрудника; такой файл не нужен")
    if expected is not None and ids != set(expected):
        raise ValueError("Состав Excel не совпал с полной выборкой заказов кабинета")
    employee_id, employee = next(iter(found_employees.items()))
    if expected:
        latest = max(expected.values(), key=lambda o: o["created_at"])
        employee = dict(employee, source_name=latest["source_name"], phone_sha256=latest["phone_sha256"],
                        name=display_name(latest["source_name"], employee["user_id"], latest["phone_sha256"], aliases))
    return {"name": employee["name"], "employee_id": employee_id,
            "user_id": employee.get("user_id"), "phone_sha256": employee["phone_sha256"],
            "source_name": employee["source_name"], "orders": checked_orders,
            "identity_basis": ("cabinet_user_id" if expected is not None else "configured_user_id") if employee.get("user_id") else "xlsx_phone_sha256", "period": [start.isoformat(), end.isoformat()],
            "litres": str(sum((Decimal(o["litres"]) for o in operations), Decimal(0))),
            "amount": str(total), "order_count": len(ids), "operations": operations,
            "active": any(Decimal(o["litres"]) > 0 for o in operations)}


def formatted(value):
    return f"{Decimal(value).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):,.2f}".replace(",", " ").replace(".", ",")


def caption(report):
    start, end = [date.fromisoformat(v) for v in report["period"]]
    return (f"{start:%d.%m} - {end:%d.%m}\nЯндекс заправки\n"
            f"на склад ({report['name']}) {formatted(report['litres'])} л на сумму закупки {formatted(report['amount'])} рублей")


def order_fingerprint(report):
    return hashlib.sha256("\n".join(sorted(o["id"] for o in report["orders"])).encode()).hexdigest()


def render_yandex_report(reports, run_date):
    reports = label_reports([r for r in reports if r["active"]])
    return f"# Яндекс Заправки — {run_date:%d.%m.%Y}\n\n" + "\n\n".join(f"### {r['label']}\n<!-- yandex-employee:{r['employee_id']} phone:{r['phone_sha256']} orders:{order_fingerprint(r)} -->\n" + caption(r) for r in reports) + "\n"


def plain_report(result):
    return re.sub(r"<!-- yandex-employee:[^\n]+ -->\n", "", result)


def write_outputs(reports, root, run_date):
    reports = label_reports([r for r in reports if r["active"]])
    result = render_yandex_report(reports, run_date)
    for suffix in ("md", "txt"):
        content = plain_report(result) if suffix == "txt" else result
        (root / f"Яндекс — {run_date:%d.%m.%Y}.{suffix}").write_text(content, encoding="utf-8")
    book = Workbook()
    sheet = book.active
    sheet.title = "На склад"
    sheet.append(["Сотрудник", "С", "По", "Литры", "Сумма закупки, руб.", "Идентификатор сотрудника"])
    for report in reports:
        sheet.append([report["name"], *report["period"], float(report["litres"]), float(report["amount"]), report["employee_id"]])
        sheet.cell(sheet.max_row, 1).data_type = "s"
        for column in (4, 5):
            sheet.cell(sheet.max_row, column).number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    for column, width in {"A": 22, "B": 16, "C": 16, "D": 18, "E": 24, "F": 44}.items():
        sheet.column_dimensions[column].width = width
    book.save(root / f"Яндекс — {run_date:%d.%m.%Y}.xlsx")
    return result

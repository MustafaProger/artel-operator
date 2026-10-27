"""Full-month GloPro customer charges, separate from activation and balances."""
from __future__ import annotations

from calendar import monthrange
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .calculator import CalculationError, DEFAULT_RULES, _date, _decimal, _ru, company_identity
from .ordering import alphabet_key


STATUS_LABELS = {"ready": "Есть операции", "no_data": "Нет операций",
                 "fully_reversed": "Операции полностью возвращены"}
MEANING = ("К оплате нам — стоимость операций за месяц с учётом клиентских скидок, "
           "наценок и возвратов. Платежи и начальный остаток не вычитаются; "
           "это начисление за период, а не остаток задолженности.")
ROUNDING = ("По каждой фирме сумма округлена до копеек, литраж — до 0,001 л. "
            "Общий итог равен сумме строк фирм. В детализации разница округления "
            "показана отдельной строкой, если она есть.")


def _text(value):
    # Remove control characters that cannot be serialized into XLSX/XML.
    return re.sub(r"[\x00-\x1f\x7f]", " ", str(value or "")).strip()


def _markdown(value):
    value = _text(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    for char in ("\\", "|", "*", "_", "[", "]", "`", "#"):
        value = value.replace(char, "\\" + char)
    return value


def _rounded(value, key):
    places = Decimal("0.001") if key == "litres" else Decimal("0.01")
    result = value.quantize(places, rounding=ROUND_HALF_UP)
    return abs(result) if result == 0 else result


def _amounts(item):
    return {key: _decimal(item.get(key), key) for key in ("litres", "customer_total")}


def _period(record):
    period = [record.get("period_start"), record.get("period_end")]
    if not all(period):
        period = record.get("period", [])
    if len(period) != 2 or not all(period):
        raise CalculationError("Не указан период месячного отчёта")
    first, last = (_date(value, "период месячного отчёта") for value in period)
    if first.day != 1 or (first.year, first.month) != (last.year, last.month) or last.day != monthrange(first.year, first.month)[1]:
        raise CalculationError("Месячный отчёт требует полный календарный месяц")
    return first, last


def _identity(item):
    name = _text(item.get("client") or item.get("name"))
    ident = str(item.get("client_id") or item.get("id") or "")
    if not name:
        raise CalculationError("Не указано название фирмы в месячном отчёте")
    return name, ident


def _matches(item, name, ident):
    # Two known portal IDs cannot be matched solely by their display names.
    if ident and item["client_id"]:
        return ident == item["client_id"]
    return company_identity(name) == company_identity(item["client"])


def _rule_text(rules):
    settings = {**DEFAULT_RULES, **(rules or {})}
    special = [_text(client) for client in settings["special_clients"]]
    multiplier = _decimal(settings["special_customer_multiplier"], "наценка клиента")
    text = ("Обычные фирмы, Китай и НК АРТЭЛЬ: «Стоимость» (P) из выгрузки GloPro; "
            "клиентская скидка или наценка уже включена, повторно не применяется.")
    if special:
        factor = format(multiplier, "f").replace(".", ",")
        text += (f" Для {', '.join(special)} по дизелю: «Стоимость ТО» (M) × {factor}, "
                 "согласно настройкам оператора. Сумма поставщику в итог «к оплате нам» не входит.")
    return text, {"special_clients": special, "special_customer_multiplier": str(multiplier)}


def _company(report, period):
    name, ident = _identity(report)
    status = report.get("status")
    if status not in STATUS_LABELS:
        raise CalculationError(f"{name}: отчёт не прошёл расчёт и проверку")
    if report.get("period") and list(report["period"]) != list(period):
        raise CalculationError(f"{name}: период фирмы не совпадает с месячным отчётом")
    audit = report.get("audit", {})
    raw = _amounts(audit.get("raw_totals") or report.get("totals", {}))
    if status in {"no_data", "fully_reversed"} and any(raw.values()):
        raise CalculationError(f"{name}: нулевой статус при ненулевых итогах")
    raw_details = audit.get("raw_groups")
    if raw_details is None:
        raw_details = audit.get("raw_fuels")
    if raw_details is None:
        raw_details = report.get("holders") if report.get("weekly_kind") == "nk_artel" else report.get("fuels", [])
    details = []
    detail_raw_totals = {key: Decimal(0) for key in raw}
    for item in sorted(raw_details, key=lambda r: (alphabet_key(r.get("holder", "")), alphabet_key(r.get("fuel", "")))):
        amounts = _amounts(item)
        for key in raw:
            detail_raw_totals[key] += amounts[key]
        details.append({"holder": _text(item.get("holder")), "fuel": _text(item.get("fuel")),
                        **{key: str(_rounded(value, key)) for key, value in amounts.items()},
                        "kind": "operations"})
    if any(abs(detail_raw_totals[key] - raw[key]) > Decimal("0.0000001") for key in raw):
        raise CalculationError(f"{name}: детализация не сходится с итогами фирмы")
    displayed = {key: _rounded(value, key) for key, value in raw.items()}
    adjustment = {key: displayed[key] - sum((Decimal(row[key]) for row in details), Decimal(0)) for key in raw}
    if any(adjustment.values()):
        details.append({"holder": "", "fuel": "Корректировка округления", "kind": "rounding",
                        **{key: str(value) for key, value in adjustment.items()}})
    if not details:
        details.append({"holder": "", "fuel": "—", "litres": "0.000", "customer_total": "0.00", "kind": "operations"})
    return {"client": name, "client_id": ident, "status": status,
            **{key: str(value) for key, value in displayed.items()}, "details": details,
            "raw_totals": {key: str(value) for key, value in raw.items()}}


def _zero_company(name, ident):
    return {"client": name, "client_id": ident, "status": "no_data",
            "totals": {"litres": "0", "customer_total": "0"}, "fuels": []}


def _build(reports, period, empty_clients, clients):
    companies = []
    for report in reports:
        company = _company(report, period)
        if any(_matches(item, company["client"], company["client_id"]) or company_identity(item["client"]) == company_identity(company["client"]) for item in companies):
            raise CalculationError("Повтор фирмы в месячном отчёте: сначала объедините проверенные договоры")
        companies.append(company)
    for empty in empty_clients:
        name, ident = _identity(empty)
        if empty.get("reason") not in (None, "no_operations"):
            raise CalculationError(f"{name}: отсутствие выгрузки не подтверждает отсутствие операций")
        if not any(_matches(item, name, ident) for item in companies):
            companies.append(_company(_zero_company(name, ident), period))
    for client in clients:
        name, ident = _identity(client)
        if not any(_matches(item, name, ident) for item in companies):
            raise CalculationError(f"{name}: нет проверенного отчёта или подтверждения отсутствия операций")
    return sorted(companies, key=lambda r: alphabet_key(r["client"]))


def _render(companies, totals, first, last, rule_text):
    lines = [f"# Месячный отчёт GloPro — {first:%m.%Y}", "",
             f"Период: {first:%d.%m.%Y}–{last:%d.%m.%Y}, включительно (Europe/Moscow).", "",
             MEANING, "", rule_text, "", ROUNDING, "",
             f"Проверено фирм: {len(companies)}; с операциями: {sum(r['status'] == 'ready' for r in companies)}; "
             f"без операций: {sum(r['status'] == 'no_data' for r in companies)}; "
             f"полностью возвращены: {sum(r['status'] == 'fully_reversed' for r in companies)}.", "",
             "| Фирма | Литры | К оплате нам, руб. | Статус |", "| --- | ---: | ---: | --- |"]
    for company in companies:
        lines.append(f"| {_markdown(company['client'])} | {_ru(company['litres'], 3)} | {_ru(company['customer_total'], 2)} | {STATUS_LABELS[company['status']]} |")
    lines.extend([f"| **ИТОГО** | **{_ru(totals['litres'], 3)}** | **{_ru(totals['customer_total'], 2)}** | |", "", "## Детализация", ""])
    for company in companies:
        lines.extend([f"### {_markdown(company['client'])}", "", "| Держатель | Топливо | Литры | К оплате нам, руб. |", "| --- | --- | ---: | ---: |"])
        for detail in company["details"]:
            lines.append(f"| {_markdown(detail['holder']) or '—'} | {_markdown(detail['fuel'])} | {_ru(detail['litres'], 3)} | {_ru(detail['customer_total'], 2)} |")
        lines.extend([f"| **Итого** | | **{_ru(company['litres'], 3)}** | **{_ru(company['customer_total'], 2)}** |", ""])
    return "\n".join(lines)


def _append(sheet, values):
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            # Explicit string cell type prevents =, +, -, @ input becoming formulas.
            cell.data_type = "s"


def _style_table(sheet, widths, number_columns):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="246B54")
    for row in sheet.iter_rows(min_row=2):
        for index, places in number_columns.items():
            row[index - 1].number_format = "#,##0." + "0" * places
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _xlsx(companies, totals, first, last, rule_text, path):
    book = Workbook()
    summary = book.active
    summary.title = "Сводка по фирмам"
    _append(summary, ["Фирма", "Литры", "К оплате нам, руб.", "Статус"])
    detail = book.create_sheet("Детализация")
    _append(detail, ["Фирма", "Держатель", "Топливо", "Литры", "К оплате нам, руб.", "Статус"])
    for company in companies:
        _append(summary, [company["client"], Decimal(company["litres"]), Decimal(company["customer_total"]), STATUS_LABELS[company["status"]]])
        for item in company["details"]:
            _append(detail, [company["client"], item["holder"], item["fuel"], Decimal(item["litres"]), Decimal(item["customer_total"]), STATUS_LABELS[company["status"]]])
    _append(summary, ["ИТОГО", Decimal(totals["litres"]), Decimal(totals["customer_total"]), ""])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    _style_table(summary, {"A": 44, "B": 18, "C": 24, "D": 36}, {2: 3, 3: 2})
    _style_table(detail, {"A": 44, "B": 30, "C": 32, "D": 18, "E": 24, "F": 36}, {4: 3, 5: 2})
    notes = book.create_sheet("Период и правила")
    for values in (["Параметр", "Значение"], ["Начало периода", first.isoformat()], ["Конец периода", last.isoformat()],
                   ["Часовой пояс", "Europe/Moscow"], ["Смысл суммы", MEANING], ["Правила расчёта", rule_text], ["Округление", ROUNDING]):
        _append(notes, values)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    book.save(path)
    book.close()


def write_monthly_outputs(reports, record, root, rules=None, *, empty_clients=(), clients=()):
    """Write verified monthly Markdown/TXT/XLSX and return JSON-safe results.

    ``reports`` already contain one calculated/merged report per firm. ``clients``
    is the provider manifest; an uncovered firm must never silently become zero.
    Contract-level coverage is verified by the caller before invoking this writer.
    ``empty_clients`` contains only independently verified no-operation firms.
    """
    first, last = _period(record)
    period = [first.isoformat(), last.isoformat()]
    with localcontext() as context:
        context.prec = 40
        companies = _build(reports, period, empty_clients, clients)
        totals = {key: str(sum((Decimal(company[key]) for company in companies), _rounded(Decimal(0), key))) for key in ("litres", "customer_total")}
        raw_totals = {key: str(sum((Decimal(company["raw_totals"][key]) for company in companies), Decimal(0))) for key in totals}
        rounding = {key: str(Decimal(totals[key]) - Decimal(raw_totals[key])) for key in totals}
    rule_text, applied_rules = _rule_text(rules)
    report_text = _render(companies, totals, first, last, rule_text)
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    stem = f"Месячный отчёт GloPro — {first:%Y-%m}"
    files = [root / f"{stem}.{extension}" for extension in ("md", "txt", "xlsx")]
    _xlsx(companies, totals, first, last, rule_text, files[2])
    for target in files[:2]:
        target.write_text(report_text, encoding="utf-8")
    public_companies = [{key: company[key] for key in ("client", "client_id", "status", "litres", "customer_total")} for company in companies]
    return {"report": report_text, "files": [str(path) for path in files],
            "metrics": {"client_count": len(companies), "active_client_count": sum(r["status"] == "ready" for r in companies),
                        "no_data_client_count": sum(r["status"] == "no_data" for r in companies),
                        "fully_reversed_client_count": sum(r["status"] == "fully_reversed" for r in companies),
                        "total_litres": totals["litres"], "total_customer_amount": totals["customer_total"],
                        "companies": public_companies, "output_period": period},
            "audit": {"monthly_summary": {"period": period, "companies": companies, "totals": totals,
                                           "raw_totals": raw_totals, "rounding_adjustment": rounding,
                                           "rules": applied_rules, "amount_meaning": MEANING}}}

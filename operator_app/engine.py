from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from threading import Lock
from urllib.parse import quote
from zoneinfo import ZoneInfo
import hashlib
import json
import re
import shutil
import sqlite3
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from . import storage
from .calculator import calculate_workbook, company_identity, merge_reports, render_report, workbook_client
from .clients import client_is_excluded
from .config import DATA, DAYS, get_operator, load_operators, latest_run_date, period_for, client_period_for, schedule_matches, seo_slot
from .publication import save_report_sections
from .ordering import alphabet_key

POOL = ThreadPoolExecutor(max_workers=1, thread_name_prefix="operator")
RUN_LOCK = Lock()
SUBMIT_LOCK = Lock()
HANDLERS = {}
PROCESSORS = {}


@dataclass
class ProcessResult:
    """Operator-specific result; executor owns persistence, source hashes and ZIP."""
    status: str
    report: str = ""
    error: str | None = None
    metrics: dict = field(default_factory=dict)
    audit: dict = field(default_factory=dict)


def glopro_download(conf, record, directory, progress):
    from .glopro import GloProConnector
    creds = storage.credentials()
    if not creds:
        raise ValueError("Подключите GloPro на странице «Подключение»")
    def connected_progress(event):
        if event.get("stage") == "clients":
            storage.set_setting("connection_verified", True)
        progress(event)
    return GloProConnector(*creds).download_reports(
        record["period_start"], record["period_end"], directory,
        clients=conf.get("clients", []), progress=connected_progress,
        excluded_clients=conf.get("excluded_clients", []), run_date=record["run_date"])


HANDLERS["glopro"] = glopro_download


def yandex_download(conf, record, directory, progress):
    from .yandex import download_reports
    return download_reports(conf, record, directory, progress)


HANDLERS["yandex"] = yandex_download


def seo_download(conf, record, directory, progress):
    from .seo import prepare_article
    return prepare_article(conf, record, directory, progress)


HANDLERS["seo"] = seo_download


def submit(ident="glopro", run_date=None, trigger="manual", uploads=None, *, edition_key=None):
    conf = get_operator(ident)
    if edition_key is not None and (conf["kind"] != "seo" or trigger != "manual" or run_date is not None
                                    or not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,79}", edition_key)):
        raise ValueError("Внеплановый выпуск: нужен SEO-оператор и постоянный edition_key без run_date")
    if conf["kind"] not in HANDLERS or conf["kind"] not in PROCESSORS:
        raise ValueError("Для этого оператора ещё не подключены получение и обработка результатов")
    scheduled_date = date.fromisoformat(run_date) if run_date else latest_run_date()
    if conf["kind"] == "seo":
        if uploads is not None:
            raise ValueError("SEO: импорт XLSX не поддерживается")
        today = datetime.now(ZoneInfo(conf["schedule"]["timezone"])).date()
        scheduled_date = seo_slot(conf, date.fromisoformat(run_date) if run_date else today)
        if scheduled_date != seo_slot(conf, today):
            raise ValueError("SEO: допускается только актуальный выпуск; старые статьи не публикуются пачкой")
        from .seo import _get as seo_edition, next_due as seo_next_due
        due = seo_next_due({**conf, "enabled": True})
        if edition_key is None and not seo_edition(conf["id"], str(scheduled_date)) and due and due.date() > today:
            raise ValueError("SEO: предыдущая статья уже опубликована; следующий выпуск " + due.date().isoformat())
    if not run_date and conf["kind"] == "yandex":
        while scheduled_date.weekday() != 1:
            scheduled_date -= timedelta(days=1)
    if scheduled_date > datetime.now(ZoneInfo(conf["schedule"]["timezone"])).date():
        raise ValueError("Нельзя рассчитывать ещё не завершившийся период")
    if conf["kind"] == "yandex":
        from .yandex_reports import period_for as yandex_period
        start, end = yandex_period(scheduled_date)
    elif conf["kind"] == "seo":
        start = end = scheduled_date
    else:
        start, end = period_for(scheduled_date)
    with SUBMIT_LOCK:
        record = storage.create_run(ident, scheduled_date, start, end, trigger,
                                    require_idle=True, edition_key=edition_key)
        # Snapshot now: edits made during a queued/running job apply next time.
        try:
            POOL.submit(execute, conf, record, uploads)
        except Exception:
            # A rejected dispatch must not leave a permanent queued reservation.
            record.update(status="failed", finished_at=storage.now_iso(),
                          error="Не удалось передать запуск исполнителю. Доступен ручной повтор.")
            storage.save_run(record)
            raise
    return record


def _file(record, path, root):
    rel = path.relative_to(root).as_posix()
    return dict(name=rel, url=f"/api/runs/{record['id']}/files/{quote(rel)}", size=path.stat().st_size,
                sha256=hashlib.sha256(path.read_bytes()).hexdigest())


def _summary_xlsx(reports, target):
    book = Workbook()
    ws = book.active
    ws.title = "Активация"
    ws.append(["Клиент", "Топливо", "Литры", "Сумма клиента, руб", "Цена поставщика, руб/л", "Держатель", "Для поставщика, руб"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="246B54")
    for report in sorted(reports, key=lambda item: alphabet_key(item["client"])):
        if report.get("weekly_kind"):
            entries = (report["holders"] if report["weekly_kind"] == "nk_artel" else
                       [{"fuel": "", "holder": "", **report["totals"]}])
            for entry in entries:
                ws.append([report["client"], entry["fuel"], float(entry["litres"]), float(entry["customer_total"]), None,
                           entry["holder"], float(report["supplier_total"]) if report["weekly_kind"] == "china" else None])
                for col in (1, 2, 6):
                    ws.cell(ws.max_row, col).data_type = "s"
                ws.cell(ws.max_row, 3).number_format = "#,##0.00" if report["weekly_kind"] == "nk_artel" else "#,##0.000"
                ws.cell(ws.max_row, 4).number_format = "#,##0.00"
                ws.cell(ws.max_row, 7).number_format = "#,##0.000"
            continue
        for fuel in report["fuels"]:
            # Never allow client-controlled cell strings to become Excel formulas.
            client = report["client"]
            if client.startswith(("=", "+", "-", "@")):
                client = "'" + client
            ws.append([client, fuel["fuel"], float(fuel["litres"]), float(fuel["customer_total"]), float(fuel["supplier_price"])])
            for col in (3, 4, 5):
                ws.cell(ws.max_row, col).number_format = "#,##0.000" if col == 3 else "#,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {"A": 42, "B": 22, "C": 17, "D": 25, "E": 29, "F": 32, "G": 25}.items():
        ws.column_dimensions[col].width = width
    book.save(target)


def process_glopro(conf, record, sources, root, progress, *, imported=False):
    """Calculate GloPro sources and write its domain-specific result files."""
    result = ProcessResult(status="failed")
    reports, failures = [], []
    excluded = {}
    schedule_skipped = {}

    def remember_schedule_skip(client):
        name = client.get("client") or client.get("name") or ""
        ident = str(client.get("client_id") or client.get("id") or "")
        schedule_skipped[ident or company_identity(name)] = {"client": name, "client_id": ident, "reason": "tuesday_only"}

    def remember_exclusion(client):
        name = client.get("client") or client.get("name") or ""
        ident = str(client.get("client_id") or client.get("id") or "")
        key = ("id", ident) if ident else ("name", company_identity(name))
        excluded[key] = {"client": name, "client_id": ident, "reason": "excluded_client"}

    for event in record.get("events", []):
        if event.get("stage") == "skipped" and event.get("reason") == "excluded_client":
            remember_exclusion(event)
        elif event.get("stage") == "skipped" and event.get("reason") == "tuesday_only":
            remember_schedule_skip(event)
    for source in sources:
        if client_is_excluded(source, conf.get("excluded_clients", [])):
            remember_exclusion(source)
            progress(dict(stage="skipped", reason="excluded_client", client=source.get("client", ""),
                          client_id=source.get("client_id", ""), message=f"Исключено из расчёта: {source.get('client', '')}"))
            continue
        try:
            client = source.get("client") or workbook_client(source["path"])
            period = client_period_for(date.fromisoformat(record["run_date"]), client, source.get("client_id"))
            if period is None:
                skipped = {**source, "client": client}
                remember_schedule_skip(skipped)
                progress(dict(stage="skipped", reason="tuesday_only", client=client,
                              client_id=source.get("client_id", ""), message=f"Только вторничный запуск: {client}"))
                continue
            rules = {**conf.get("rules", {}), "supplier": conf.get("supplier", "Новое имя"),
                     "date_from": period[0], "date_to": period[1], "run_date": record["run_date"],
                     "client": client, "client_id": source.get("client_id")}
            report = calculate_workbook(source["path"], rules)
            source["client"] = report["client"]
            report["period"] = [period[0].isoformat(), period[1].isoformat()]
            if source.get("contract_id"):
                report["contract_id"] = str(source["contract_id"])
            if source.get("client_id"):
                report["client_id"] = str(source["client_id"])
            if client_is_excluded(report, conf.get("excluded_clients", [])):
                remember_exclusion(report)
                progress(dict(stage="skipped", reason="excluded_client", client=report["client"],
                              client_id=report.get("client_id", ""), message=f"Исключено из расчёта: {report['client']}"))
                continue
            reports.append(report)
            progress(dict(stage="calculated", message=f"Проверено: {report['client']}"))
        except ValueError as exc:
            failures.append({"file": Path(source["path"]).name, "error": str(exc)})
    calculated_sources = list(reports)
    grouped = {}
    for report in reports:
        identity = ("client_id", report["client_id"]) if not imported and report.get("client_id") else ("name", company_identity(report["client"]))
        grouped.setdefault(identity, []).append(report)
    reports = []
    for group in grouped.values():
        if len(group) == 1:
            reports.append(group[0])
            continue
        contracts = [r.get("contract_id") for r in group]
        client_ids = {r.get("client_id") for r in group}
        if imported or None in client_ids or len(client_ids) != 1 or any(c is None for c in contracts) or len(contracts) != len(set(contracts)):
            raise ValueError("Повтор фирмы в исходных файлах: нельзя складывать пересекающиеся выгрузки")
        try:
            merged = merge_reports(group, conf.get("rules", {}))
            merged["period"] = group[0]["period"]
            reports.append(merged)
        except ValueError as exc:
            failures.append({"file": group[0]["client"], "error": str(exc)})
    reports.sort(key=lambda item: alphabet_key(item["client"]))
    result.metrics["client_count"] = len(reports)
    result.metrics["active_client_count"] = sum(report.get("status") == "ready" for report in reports)
    result.metrics["source_count"] = len(sources)
    result.metrics["failures"] = failures
    result.metrics["excluded_clients"] = list(excluded.values())
    result.metrics["excluded_client_count"] = len(excluded)
    result.metrics["schedule_skipped_clients"] = list(schedule_skipped.values())
    result.metrics["empty_clients"] = [event for event in record.get("events", []) if event.get("reason") == "no_operations"]
    company_periods = [{"client": r["client"], "period": r["period"]} for r in reports]
    result.metrics["company_periods"] = company_periods
    if company_periods:
        result.metrics["output_period"] = [min(r["period"][0] for r in reports), max(r["period"][1] for r in reports)]
    identities = [company_identity(r["client"]) for r in reports]
    if len(identities) != len(set(identities)):
        raise ValueError("Повтор фирмы в исходных файлах: нельзя складывать пересекающиеся выгрузки")
    if failures:
        result.status = "needs_review"
        result.error = f"Нужна проверка {len(failures)} из {len(sources)} файлов. Итоговая активация не выпущена."
        (root / "Ошибки проверки.json").write_text(json.dumps(failures, ensure_ascii=False, indent=2), encoding="utf-8")
    elif not any(report.get("status") == "ready" or report.get("weekly_kind") for report in reports):
        result.status = "no_data"
        result.error = ("Все переданные фирмы временно исключены из расчёта." if excluded and not reports else
                        "Фирмы проверены: за выбранный период заправок нет. XLSX не запрашивались." if not sources else
                        "За выбранный период нет операций для активации. Исходные файлы проверены и сохранены.")
    else:
        report_text = render_report([r for r in reports if r.get("status") == "ready" or r.get("weekly_kind")], date.fromisoformat(record["run_date"]))
        label = date.fromisoformat(record["run_date"]).strftime("%d.%m.%Y")
        (root / f"Активация — {label}.md").write_text(report_text, encoding="utf-8")
        (root / f"Активация — {label}.txt").write_text(report_text, encoding="utf-8")
        _summary_xlsx(reports, root / f"Активация — {label}.xlsx")
        result.report = report_text
        result.status = "completed" if any(r["status"] == "ready" for r in reports) else "no_data"
        export_dir = conf.get("obsidian_output")
        if export_dir:
            directory = Path(export_dir).expanduser()
            directory.mkdir(parents=True, exist_ok=True)
            destination = directory / f"Активация — {label}.md"
            save_report_sections(destination, report_text)
    result.audit = {"reports": reports, "calculated_sources": calculated_sources, "errors": failures,
                    "empty_clients": result.metrics["empty_clients"],
                    "excluded_clients": list(excluded.values()), "schedule_skipped_clients": list(schedule_skipped.values()),
                    "company_periods": company_periods}
    return result


PROCESSORS["glopro"] = process_glopro


def process_yandex(conf, record, sources, root, progress, *, imported):
    from .yandex_reports import read_report, write_outputs, label_reports, render_yandex_report, plain_report, period_for as yandex_period
    from .yandex_publication import save_yandex_sections, merge_yandex_sections
    start, end = yandex_period(date.fromisoformat(record["run_date"]))
    aliases = conf.get("employee_aliases", conf.get("employees", []))
    manifests = [e["manifest"] for e in record["events"] if e.get("stage") == "discovered" and "manifest" in e]
    manifest = manifests[0] if len(manifests) == 1 else None
    reports, failures = [], []
    expected_by_user = {}
    if manifest:
        for order in manifest["orders"]:
            expected_by_user.setdefault(order["user_id"], []).append(order)
    for source in sources:
        try:
            expected = None if imported else expected_by_user.get(source.get("user_id"), [])
            report = read_report(source["path"], start, end, conf["company"], aliases=aliases, expected_orders=expected)
            if not imported and report["user_id"] != source.get("user_id"):
                raise ValueError("ID сотрудника Excel не совпал с источником")
            report["report_id"] = source.get("report_id")
            reports.append(report)
        except ValueError as exc:
            failures.append({"file": Path(source["path"]).name, "error": str(exc)})
    if len({r["employee_id"] for r in reports}) != len(reports):
        raise ValueError("Повтор сотрудника Яндекса: нельзя складывать пересекающиеся файлы")
    all_ids = [o["id"] for r in reports for o in r["orders"]]
    if len(set(all_ids)) != len(all_ids):
        raise ValueError("Один заказ Яндекса попал в несколько файлов")
    empty = [e for e in record["events"] if e.get("reason") == "no_operations"]
    if not imported:
        completed = [e for e in record["events"] if e.get("activity_check_complete")]
        active_users = {uid for uid, orders in expected_by_user.items() if any(o["active"] for o in orders)}
        empty_users = set(expected_by_user) - active_users
        if (not manifest or manifest.get("complete") is not True
                or manifest.get("period") != [str(start), str(end)] or manifest.get("timezone") != "UTC+3"
                or manifest.get("order_count") != len(manifest["orders"])
                or len({o["id"] for o in manifest["orders"]}) != len(manifest["orders"])
                or len(completed) != 1 or completed[0].get("order_count") != manifest["order_count"]
                or set(completed[0].get("employee_ids", [])) != set(expected_by_user)
                or {r["user_id"] for r in reports} != active_users
                or {e.get("user_id") for e in empty if e.get("user_id")} != empty_users):
            failures.append({"error": "Не подтверждена полная выборка недели или покрытие всех сотрудников по ID"})
    metrics = {"client_count": len(reports), "active_client_count": sum(r["active"] for r in reports),
               "source_count": len(sources), "empty_clients": empty, "failures": failures,
               "company_periods": [{"client": r["name"], "user_id": r["user_id"], "period": r["period"]} for r in reports]}
    audit = {"reports": reports, "errors": failures, "empty_clients": empty, "imported": imported,
             "cabinet_manifest": manifest,
             "coverage": "imported_files_only" if imported else "all_week_orders",
             "litre_reconciliation": "cabinet litres rounded per row to XLSX precision 0.01, HALF_UP"}
    if failures:
        return ProcessResult("needs_review", error="Отчёты Яндекса не прошли проверку; итог не опубликован.", metrics=metrics, audit=audit)
    # Do not use display names as identities or let duplicate names overwrite a source.
    label_reports(reports)
    renames = [(Path(source["path"]), Path(source["path"]).parent / f"Яндекс. {report['label']}.xlsx")
               for source, report in zip(sources, reports)]
    for path, destination in renames:
        if path != destination and destination.exists():
            raise ValueError("Имя файла Яндекса уже занято: повторная выгрузка сотрудника")
    for source, (path, destination) in zip(sources, renames):
        if path != destination:
            path.rename(destination)
            source["path"] = str(destination)
    reports.sort(key=lambda r: (alphabet_key(r["name"]), r["employee_id"]))
    if not any(r["active"] for r in reports):
        return ProcessResult("no_data", error="Подтверждено отсутствие заправок за неделю.", metrics=metrics, audit=audit)
    report_text = render_yandex_report(reports, date.fromisoformat(record["run_date"]))
    if conf.get("obsidian_output"):
        destination = Path(conf["obsidian_output"]).expanduser()
        destination.mkdir(parents=True, exist_ok=True)
        target = destination / f"Яндекс Заправки — {date.fromisoformat(record['run_date']):%d.%m.%Y}.md"
        try:
            merge_yandex_sections(target.read_text(encoding="utf-8") if target.exists() else "", report_text)
        except ValueError as exc:
            failures.append({"error": str(exc)})
            return ProcessResult("needs_review", error=str(exc), metrics=metrics, audit=audit)
        save_yandex_sections(target, report_text)
    write_outputs(reports, root, date.fromisoformat(record["run_date"]))
    return ProcessResult("completed", report=plain_report(report_text), metrics=metrics, audit=audit)


PROCESSORS["yandex"] = process_yandex


def process_seo(conf, record, sources, root, progress, *, imported=False):
    from .seo import publish_article
    result = publish_article(conf, record, sources, root, progress)
    return ProcessResult(**result)


PROCESSORS["seo"] = process_seo


def execute(conf, record, uploads=None):
    root = DATA / "runs" / record["id"]
    originals = root / ("Файлы по фирмам" if conf["kind"] == "glopro" else "Исходные файлы")

    def progress(event):
        record["events"].append({"at": storage.now_iso(), **event})
        storage.save_run(record)

    with RUN_LOCK:
        try:
            originals.mkdir(parents=True, exist_ok=True)
            record["status"] = "running"
            storage.save_run(record)
            (root / "Инструкция.md").write_text(conf["markdown"], encoding="utf-8")
            if uploads is not None:
                sources = []
                for upload in uploads:
                    destination = originals / upload["name"]
                    shutil.copy2(upload["path"], destination)
                    sources.append({"path": str(destination)})
            else:
                sources = HANDLERS[conf["kind"]](conf, record, originals, progress)
            verified_empty = (conf["kind"] in {"glopro", "yandex"} and uploads is None
                              and any(e.get("activity_check_complete") for e in record["events"])
                              and any(e.get("reason") == "no_operations" for e in record["events"]))
            if not sources and not verified_empty:
                raise ValueError("Не получено ни одного файла. Пустая выгрузка не считается готовым комплектом.")
            result = PROCESSORS[conf["kind"]](conf, record, sources, root, progress, imported=uploads is not None)
            if not isinstance(result, ProcessResult) or result.status not in {"completed", "needs_review", "no_data"}:
                raise ValueError("Обработчик вернул некорректный результат")
            reserved = {"id", "operator_id", "run_date", "period_start", "period_end", "trigger", "created_at", "events", "files", "finished_at", "status", "report", "error"}
            if reserved.intersection(result.metrics):
                raise ValueError("Метрики обработчика не могут изменять сведения о запуске")
            record.update(result.metrics)
            record.update(status=result.status, report=result.report, error=result.error)
            audit = {**result.audit, "run_id": record["id"], "activation_date": record["run_date"],
                     "period": [record["period_start"], record["period_end"]],
                     "operator_sha256": hashlib.sha256(conf["markdown"].encode()).hexdigest(),
                     "sources": [{"name": Path(s["path"]).name, "sha256": hashlib.sha256(Path(s["path"]).read_bytes()).hexdigest()} for s in sources]}
            if record.get("reused_from_run"):
                audit["reused_from_run"] = record["reused_from_run"]
            audit_name = "Проверка SEO.json" if conf["kind"] == "seo" else "Проверка расчётов.json"
            (root / audit_name).write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
            label = {"completed": "Готовый комплект", "no_data": "Нет операций"}.get(record["status"], "На проверку")
            output_period = record.get("output_period", [record["period_start"], record["period_end"]])
            archive = root / f"{label} — {output_period[0]} — {output_period[1]}.zip"
            with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
                for path in sorted(root.rglob("*"), key=(lambda p: alphabet_key(p.relative_to(root))) if conf["kind"] == "yandex" else None):
                    if path.is_file() and path != archive:
                        bundle.write(path, path.relative_to(root))
        except Exception as exc:
            record["status"] = "failed"
            # Connector errors must already be scrubbed; never serialize tracebacks/credentials.
            record["error"] = str(exc) if isinstance(exc, ValueError) else "Сбой выполнения. Проверьте подключение и исходные файлы; доступен повтор."
        finally:
            record["finished_at"] = storage.now_iso()
            try:
                record["files"] = [_file(record, p, root) for p in sorted(root.rglob("*"), key=(lambda p: alphabet_key(p.relative_to(root))) if conf["kind"] == "yandex" else None) if p.is_file()]
            except OSError:
                # Failure to read an artifact must not strand a running job.
                record.update(status="failed", files=[],
                              error="Не удалось проверить файлы результата. Проверьте доступ к каталогу запуска; доступен ручной повтор.")
            try:
                storage.save_run(record)
            finally:
                if uploads:
                    for upload in uploads:
                        Path(upload["path"]).unlink(missing_ok=True)


def tick(now=None):
    """Persistent once-per-date schedule; catch up last 7 days after a restart."""
    for conf in load_operators():
        if not conf["enabled"]:
            continue
        if conf["kind"] == "glopro" and not storage.credentials():
            continue
        if conf["kind"] == "yandex":
            from .yandex_connection import status as yandex_status
            if not yandex_status()["configured"] or yandex_status()["connecting"]:
                continue
        zone = ZoneInfo(conf["schedule"]["timezone"])
        local = (now or datetime.now(zone)).astimezone(zone)
        enabled_since = storage.setting(f"enabled_since:{conf['id']}")
        if not enabled_since:
            storage.set_setting(f"enabled_since:{conf['id']}", local.isoformat())
            continue
        since = datetime.fromisoformat(enabled_since)
        if conf["kind"] == "seo":
            # One current edition after sleep/restart, never a burst of missed posts.
            from .seo import next_due as seo_next_due
            due = seo_next_due(conf, local)
            if due and local < due:
                continue
            try:
                day = seo_slot(conf, local.date())
            except ValueError:
                continue
            scheduled = datetime.combine(day, time.fromisoformat(conf["schedule"]["time"]), zone)
            if scheduled > local:
                continue
            try:
                submit(conf["id"], day.isoformat(), trigger="schedule")
            except sqlite3.IntegrityError:
                pass
            except ValueError:
                return
            continue
        for delta in range(7, -1, -1):
            day = local.date() - timedelta(days=delta)
            scheduled = datetime.combine(day, time.fromisoformat(conf["schedule"]["time"]), zone)
            if not schedule_matches(conf, day) or scheduled > local or scheduled < since:
                continue
            try:
                submit(conf["id"], day.isoformat(), trigger="schedule")
            except sqlite3.IntegrityError:
                continue
            except ValueError:
                return

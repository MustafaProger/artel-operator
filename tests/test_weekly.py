from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from operator_app.calculator import CalculationError, calculate_workbook, merge_reports, render_report, weekly_kind
from operator_app.config import client_period_for, latest_run_date
from operator_app.publication import update_report_sections


def weekly_file(path, client="Китай", cards=None, *, start="15.09.2026", end="21.09.2026", period=True):
    """Observed GloPro layout: card blocks, service/card/file totals, footer period."""
    cards = cards if cards is not None else [("1", "Имя из Excel", [("ДТ", "1.005", "123.45", "111.10")])]
    book = Workbook()
    ws = book.active
    ws.title = "transactions"
    ws.append([None, client])
    headers = [None] * 20
    for index, value in {3: "Дата", 7: "Операция", 8: "Услуга", 9: "Количество", 12: "Стоимость ТО", 15: "Стоимость", 18: "Скидка, %"}.items():
        headers[index] = value
    ws.append(headers)
    totals = [Decimal(0)] * 3

    def summary(label, amounts):
        row = [None] * 20
        row[1] = label
        for col, amount in zip((9, 12, 15), amounts):
            row[col] = str(amount)
        ws.append(row)

    for card, holder, operations in cards:
        ws.append([None, "Карта: " + card, None, None, None, "Держатель: " + holder])
        card_totals = [Decimal(0)] * 3
        by_fuel = {}
        for operation in operations:
            by_fuel.setdefault(operation[0], []).append(operation)
        for fuel, txs in by_fuel.items():
            ws.append([None, None, "Услуга: " + fuel])
            service_totals = [Decimal(0)] * 3
            for operation in txs:
                _, litres, supplier, customer, *timestamp = operation
                row = [None] * 20
                row[3] = timestamp[0] if timestamp else datetime.strptime(start + " 00:00:00", "%d.%m.%Y %H:%M:%S")
                row[7], row[8] = "Возврат" if Decimal(str(litres)) < 0 else "Дебет", fuel
                row[9], row[12], row[15], row[18] = str(litres), str(supplier), str(customer), "arbitrary discount ignored"
                ws.append(row)
                for i, value in enumerate((litres, supplier, customer)):
                    amount = Decimal(str(value))
                    totals[i] += amount
                    card_totals[i] += amount
                    service_totals[i] += amount
            summary("Итого по " + fuel + ":", service_totals)
        summary("Итого по карте:", card_totals)
    summary("Итого по отчету:", totals)
    ws.append([None, "Общий итог по услугам"])
    # Repeated totals must never become operations.
    summary("ДТ ЭКТО", totals)
    ws.append([None, "Отчёт:", None, None, "Транзакционный отчёт"])
    if period:
        ws.append([None, "Период:", None, None, f"с {start} 00:00:00 по {end} 23:59:59"])
    book.save(path)
    return path


def calculate(path, **rules):
    return calculate_workbook(path, {"run_date": "2026-09-22", **rules})


@pytest.mark.parametrize("client", ["Китай", 'ООО «НК АРТЭЛЬ»'])
@pytest.mark.parametrize("run_day,first,last", [
    ("2026-09-22", "2026-09-15", "2026-09-21"),
    ("2026-09-01", "2026-08-25", "2026-08-31"),
    ("2027-01-05", "2026-12-29", "2027-01-04"),
])
def test_weekly_schedule_crosses_month_and_year(client, run_day, first, last):
    assert client_period_for(date.fromisoformat(run_day), client) == (date.fromisoformat(first), date.fromisoformat(last))
    assert client_period_for(date(2026, 9, 25), client) is None


def test_other_clients_and_exact_weekly_identities():
    for client in ("Китай-2", 'АО "НК АРТЭЛЬ"', 'ООО "ТЕСТ"'):
        assert client_period_for(date(2026, 9, 22), client)[0] == date(2026, 9, 18)
        assert client_period_for(date(2026, 9, 25), client)[0] == date(2026, 9, 22)
    assert weekly_kind("Переименован", "78749") == "china"
    assert weekly_kind("Китай", "111") is None
    assert latest_run_date(datetime(2026, 9, 21, 21, 10, tzinfo=timezone.utc)) == date(2026, 9, 22)


def test_china_uses_p_not_aggregate_discount_and_rounds_only_final_supplier_sum(tmp_path):
    path = weekly_file(tmp_path / "china.xlsx", cards=[("1", "Держатель", [
        ("ДТ", "1.0004", "123.45", "111.10"),
        ("Дт", "1.0004", "123.45", "111.10"),
        ("ДТ ЭКТО", "-1.0004", "-123.45", "-111.10"),
        ("АИ-95", "1", "100", "90.04"),
    ])])
    report = calculate(path)
    assert report["card_type"] is None
    assert report["totals"]["litres"] == "2.000"
    assert report["supplier_total"] == "221.216"
    assert report["totals"]["customer_total"] == "201.14"
    assert report["audit"]["raw_totals"]["litres"] == "2.0004"
    assert {fuel["fuel"] for fuel in report["fuels"]} == {"ДТ ЭКТО", "АИ-95"}
    assert len(report["audit"]["transaction_rows"]) == 4
    assert len(report["audit"]["reconciliation"]) == 6
    text = render_report([report], "2026-09-22")
    assert text == ('22.09.2026\n\n### Китай\nПоставщик: "Новое имя"\n'
                    'на склад (Китай) 2,000 л на сумму\n\n'
                    '- для поставщика 221,216 рублей\n- клиент нам 201,14 рублей\n')
    assert "руб/л" not in text


def test_nk_groups_all_current_holders_across_cards_and_diesel_labels(tmp_path):
    path = weekly_file(tmp_path / "nk.xlsx", 'ООО "НК АРТЭЛЬ"', cards=[
        ("1", "  Новый  держатель ", [("ДТ", "1.005", "100.005", "99.005")]),
        ("2", "Новый держатель", [("Дт", "1", "80", "79.20"), ("ДТ ЭКТО", "-0.1", "-8", "-7.92"), ("АИ-95", "1.005", "75", "74.25")]),
        ("3", "Эльдар", [("ДТ", "3", "240", "237.60")]),
    ])
    report = calculate(path)
    assert report["holders"] == [
        {"holder": "Новый держатель", "fuel": "ДТ ЭКТО", "litres": "1.91", "customer_total": "170.29"},
        {"holder": "Новый держатель", "fuel": "АИ-95", "litres": "1.01", "customer_total": "74.25"},
        {"holder": "Эльдар", "fuel": "ДТ ЭКТО", "litres": "3.00", "customer_total": "237.60"},
    ]
    text = render_report([report], "2026-09-22")
    assert 'на склад (Эльдар)\n• Новый держатель - ДТ ЭКТО 1,91 л | 170,29 руб' in text
    assert '• Эльдар - ДТ ЭКТО 3,00 л | 237,60 руб' in text
    assert "руб/л" not in text and "для поставщика" not in text


@pytest.mark.parametrize("sign,expected", [(1, "0.099"), (-1, "-0.099")])
def test_identical_operations_are_counted_and_supplier_is_rounded_after_sum(tmp_path, sign, expected):
    operation = ("ДТ", sign, str(Decimal("0.05") * sign), str(Decimal("0.04") * sign))
    report = calculate(weekly_file(tmp_path / "equal.xlsx", cards=[("1", "Имя", [operation, operation])]))
    assert len(report["audit"]["transaction_rows"]) == 2
    assert report["supplier_total"] == expected


def test_reconciliation_does_not_hide_small_mismatch_with_display_rounding(tmp_path):
    path = weekly_file(tmp_path / "precision.xlsx")
    book = load_workbook(path)
    for row in book.active:
        if row[1].value == "Итого по карте:":
            row[15].value = "111.104"
    book.save(path)
    with pytest.raises(CalculationError, match="не совпадает итог"):
        calculate(path)


@pytest.mark.parametrize("change", ["card", "file", "service", "missing_card", "missing_file", "missing_holder", "positive_return"])
def test_bad_totals_holder_and_return_require_review(tmp_path, change):
    path = weekly_file(tmp_path / "bad.xlsx", 'ООО "НК АРТЭЛЬ"')
    book = load_workbook(path)
    ws = book.active
    label = {"card": "Итого по карте:", "file": "Итого по отчету:", "service": "Итого по ДТ:",
             "missing_card": "Итого по карте:", "missing_file": "Итого по отчету:"}.get(change)
    for row in ws:
        if row[1].value == label and label:
            if change.startswith("missing"):
                for cell in row:
                    cell.value = None
            else:
                row[15].value = "999999"
    if change == "missing_holder":
        ws["F3"] = "Держатель: "
    if change == "positive_return":
        ws["H5"] = "Возврат"
    book.save(path)
    with pytest.raises(CalculationError):
        calculate(path)


@pytest.mark.parametrize("start,end,period", [("18.09.2026", "21.09.2026", True), ("15.09.2026", "21.09.2026", False)])
def test_old_or_missing_workbook_period_rejected_even_when_transactions_fit(tmp_path, start, end, period):
    path = weekly_file(tmp_path / "old.xlsx", start=start, end=end, period=period)
    with pytest.raises(CalculationError, match="[Пп]ериод"):
        calculate(path)


def test_wrong_hour_inside_workbook_is_not_a_whole_day(tmp_path):
    path = weekly_file(tmp_path / "hours.xlsx")
    book = load_workbook(path)
    for row in book.active:
        for cell in row:
            if isinstance(cell.value, str) and "23:59:59" in cell.value:
                cell.value = cell.value.replace("23:59:59", "22:00:00")
    book.save(path)
    with pytest.raises(CalculationError, match="Период внутри Excel"):
        calculate(path)


@pytest.mark.parametrize("stamp,accepted", [
    ("2026-09-15 00:00:00", True), ("2026-09-21 23:59:59", True),
    ("2026-09-14T21:00:00+00:00", True),
    ("2026-09-14 23:59:59", False), ("2026-09-22 00:00:00", False),
])
def test_transaction_boundaries_are_moscow_inclusive(tmp_path, stamp, accepted):
    path = weekly_file(tmp_path / "boundary.xlsx", cards=[("1", "Имя", [("Дт", 1, 10, 9, stamp)])])
    if accepted:
        assert calculate(path)["status"] == "ready"
    else:
        with pytest.raises(CalculationError, match="вне запрошенного периода"):
            calculate(path)


def test_no_card_type_or_gasoline_markup_even_if_rules_changed(tmp_path):
    path = weekly_file(tmp_path / "no-discount.xlsx", cards=[("1", "Имя", [("АИ-95", 1, 100, 99)])])
    book = load_workbook(path)
    book.active["S2"] = None
    book.save(path)
    report = calculate(path, card_type="plastic", gasoline_multiplier="7", virtual_diesel_multiplier="2")
    assert report["supplier_total"] == "99.000"
    assert report["totals"]["customer_total"] == "99.00"


@pytest.mark.parametrize("client", ["Китай", 'ООО "НК АРТЭЛЬ"'])
def test_contracts_merge_raw_values_before_rounding(tmp_path, client):
    cards = [("1", "Общий держатель", [("ДТ", "1.004", "0.005", "0.005")])]
    a = calculate(weekly_file(tmp_path / "a.xlsx", client, cards))
    b = calculate(weekly_file(tmp_path / "b.xlsx", client, cards))
    a["contract_id"], b["contract_id"] = "a", "b"
    merged = merge_reports([a, b])
    assert merged["totals"]["customer_total"] == "0.01"
    assert merged["supplier_total"] == "0.010"
    assert merged["holders"][0]["litres"] == "2.01"


@pytest.mark.parametrize("cards,status", [([], "no_data"), ([("1", "Имя", [("Дт", 1, 10, 9), ("ДТ", -1, -10, -9)])], "fully_reversed")])
def test_empty_and_fully_reversed_weekly_reports(tmp_path, cards, status):
    report = calculate(weekly_file(tmp_path / "zero.xlsx", cards=cards))
    assert report["status"] == status
    assert report["totals"]["litres"] == "0.000"
    assert report["supplier_total"] == "0.000"


def test_publication_replaces_only_addressed_sections_and_deduplicates(tmp_path):
    incoming = render_report([calculate(weekly_file(tmp_path / "c.xlsx"))], "2026-09-22")
    other = '### ООО "ДРУГАЯ"\nРучная строка без изменений.  \n\n'
    existing = '22.09.2026\n\n### Китай\nСтарый текст\n\n' + other + '### КИТАЙ\nДубликат\n'
    updated = update_report_sections(existing, incoming)
    assert updated.count("### Китай") == 1
    assert "Старый" not in updated and "Дубликат" not in updated
    assert other in updated
    assert update_report_sections(updated, incoming) == updated


def test_appending_report_is_idempotent_and_preserves_other_sections(tmp_path):
    incoming = render_report([calculate(weekly_file(tmp_path / "c.xlsx"))], "2026-09-22")
    existing = '22.09.2026\n\n### ООО "ДРУГАЯ"\nРучной отчёт.\n'
    updated = update_report_sections(existing, incoming)
    assert existing.split("\n\n", 1)[1] in updated
    assert updated.index("### Китай") < updated.index('### ООО "ДРУГАЯ"')
    assert update_report_sections(updated, incoming) == updated

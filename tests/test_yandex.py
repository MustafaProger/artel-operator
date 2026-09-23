from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch, Mock
import json

import pytest
from openpyxl import Workbook, load_workbook

from operator_app import config, engine, storage, yandex_connection
from operator_app.yandex_reports import period_for, read_report, phone_hash, caption
from operator_app.yandex import download_employee
from test_engine import isolated_engine

START, END = date(2026, 9, 15), date(2026, 9, 21)
STAFF = [dict(name='Джибриль', source_name='Джибраил', user_id='a'*32, phone_sha256=phone_hash('71111111111')),
         dict(name='Эльдар', source_name='Эльдар', user_id='b'*32, phone_sha256=phone_hash('72222222222'))]


def source(path, staff=0):
    book=Workbook(); ws=book.active; ws.title='Отчёт'
    for key,value in [('Компания','ООО НК АРТЭЛЬ'),('Период','15.09.2026 - 21.09.2026'),('Данные указаны по часовому поясу','UTC+3'),('Количество заказов в отчётном периоде',1),('Общая стоимость с НДС',3141.90)]:
        ws.append([key,None,None,value])
    ws.append([])
    ws.append(['Дата заказа','Имя пользователя','Телефон','Идентификатор заказа','Топливо','Запрос','Залито','Статус','Скидка','Стоимость'])
    ws.append(['15.09.2026',STAFF[staff]['source_name'],'71111111111' if staff==0 else '72222222222',('a' if staff==0 else 'b')*32,'АИ-100',35.74,30,'Завершён',0,3141.9])
    ws.append([None]*8+['Итого',3141.90])
    book.save(path);return path


def modify(path, cell, value):
    book=load_workbook(path);book.active[cell]=value;book.save(path);book.close()


def test_week_and_actual_fuel(tmp_path):
    assert period_for(date(2026,9,22))==(START,END)
    assert period_for(date(2027,1,5))==(date(2026,12,29),date(2027,1,4))
    with pytest.raises(ValueError,match='вторник'): period_for(date(2026,9,25))
    report=read_report(source(tmp_path/'r.xlsx'),START,END,'ООО НК АРТЭЛЬ',STAFF)
    assert Decimal(report['litres'])==30
    assert caption(report)=='15.09 - 21.09\nЯндекс заправки\nна склад (Джибриль) 30,00 л на сумму закупки 3 141,90 рублей'


@pytest.mark.parametrize('cell,value', [('D1','Другая фирма'),('D2','01.09.2026 - 21.09.2026'),('D3','UTC+4'),('D4',2),('D5',1),('J9',1),('C8','73333333333'),('B8','Другой'),('D8','invalid'),('A8','14.09.2026'),('H8','Отменён'),('G8',None),('J8','=1+1'),('G8',-30),('G7','Запрос')])
def test_untrusted_or_inconsistent_source_is_rejected(tmp_path,cell,value):
    path=source(tmp_path/'r.xlsx');modify(path,cell,value)
    with pytest.raises(ValueError):read_report(path,START,END,'ООО НК АРТЭЛЬ',STAFF)


def yandex_config(operators, output=None):
    import yaml
    value=dict(id='yandex',name='Яндекс',kind='yandex',enabled=True,company='ООО НК АРТЭЛЬ',employees=STAFF,schedule=dict(days=['tue'],time='07:00',timezone='Europe/Moscow'))
    if output:value['obsidian_output']=str(output)
    markdown='---\n'+yaml.safe_dump(value,allow_unicode=True)+'---\nНедельный отчёт.'
    (operators/'yandex.md').write_text(markdown)
    return markdown


def test_import_has_week_files_audit_alphabetic_order_and_keeps_other_sections(isolated_engine,tmp_path,monkeypatch):
    monkeypatch.setitem(engine.HANDLERS,"yandex",lambda *a: (_ for _ in ()).throw(AssertionError("network")))
    data,operators=isolated_engine; output=tmp_path/'notes';yandex_config(operators,output)
    first=engine.submit('yandex','2026-09-22','import',[dict(path=str(source(tmp_path/'eldar.xlsx',1)),name='eldar.xlsx')])
    assert first['status']=='completed'
    second=engine.submit('yandex','2026-09-22','import',[dict(path=str(source(tmp_path/'dj.xlsx')),name='dj.xlsx')])
    assert second['status']=='completed'
    assert second['period_start']=='2026-09-15'
    note=(output/'Яндекс Заправки — 22.09.2026.md').read_text()
    assert note.count('# Яндекс Заправки — 22.09.2026') == 1
    assert note.index('### Джибриль')<note.index('### Эльдар')
    assert any(f['name'].endswith('Яндекс. Джибриль.xlsx') for f in second['files'])
    audit=json.loads((data/'runs'/second['id']/'Проверка расчётов.json').read_text())
    assert audit['reports'][0]['amount']=='3141.9'
    with pytest.raises(ValueError,match='вторник'):engine.submit('yandex','2026-09-18')


def test_bad_source_does_not_publish_and_duplicates_fail(isolated_engine,tmp_path,monkeypatch):
    monkeypatch.setitem(engine.HANDLERS,"yandex",lambda *a: (_ for _ in ()).throw(AssertionError("network")))
    _,operators=isolated_engine;output=tmp_path/'notes';yandex_config(operators,output)
    bad=source(tmp_path/'bad.xlsx');modify(bad,'J9',0)
    run=engine.submit('yandex','2026-09-22','import',[dict(path=str(bad),name=bad.name)])
    assert run['status']=='needs_review' and not output.exists()
    paths=[source(tmp_path/f'{i}.xlsx') for i in range(2)]
    run=engine.submit('yandex','2026-09-22','import',[dict(path=str(p),name=p.name) for p in paths])
    assert run['status']=='failed'


def test_empty_never_creates_or_downloads(tmp_path):
    events=[]; cabinet=Mock(); page=Mock()
    employee=dict(STAFF[0],active=False,orders=[])
    assert download_employee(cabinet,page,tmp_path,employee,START,END,'ООО НК АРТЭЛЬ',[],events.append) is None
    assert not cabinet.mock_calls and not page.mock_calls
    assert events[0]['reason']=='no_operations'


def test_private_session_permissions(tmp_path,monkeypatch):
    monkeypatch.setattr(storage,'DATA',tmp_path)
    class Context:
        def storage_state(self):return {'cookies':[], 'origins':[]}
    yandex_connection.save_session(Context())
    assert yandex_connection.session_path().stat().st_mode & 0o777 == 0o600


def test_schedule_only_tuesday_and_moscow(isolated_engine):
    _,operators=isolated_engine;markdown=yandex_config(operators)
    config.parse_operator(markdown)
    with pytest.raises(ValueError,match='только вторник'):config.parse_operator(markdown.replace('- tue','- fri'))
    with pytest.raises(ValueError,match='Europe/Moscow'):config.parse_operator(markdown.replace('Europe/Moscow','UTC'))


def test_missing_employee_coverage_blocks_automatic_result(isolated_engine,tmp_path,monkeypatch):
    _,operators=isolated_engine;yandex_config(operators)
    def partial(conf,record,directory,progress):
        path=source(directory/'source.xlsx')
        return [dict(path=str(path))]
    monkeypatch.setitem(engine.HANDLERS,'yandex',partial)
    run=engine.submit('yandex','2026-09-22')
    assert run['status']=='needs_review' and not run.get('report')


def test_yandex_scheduler_is_independent_of_glopro_login(isolated_engine,monkeypatch):
    from datetime import datetime
    from zoneinfo import ZoneInfo
    _,operators=isolated_engine;yandex_config(operators)
    calls=[]
    monkeypatch.setattr(yandex_connection,'status',lambda:dict(configured=True,connecting=False))
    monkeypatch.setattr(engine,'submit',lambda *a,**kw:calls.append((a,kw)))
    storage.set_setting('enabled_since:yandex','2026-09-22T00:00:00+03:00')
    engine.tick(datetime(2026,9,22,7,0,tzinfo=ZoneInfo('Europe/Moscow')))
    assert calls==[(('yandex','2026-09-22'),{'trigger':'schedule'})]
    calls.clear()
    storage.set_setting('enabled_since:yandex','2026-09-25T00:00:00+03:00')
    engine.tick(datetime(2026,9,25,7,0,tzinfo=ZoneInfo('Europe/Moscow')))
    assert not calls

from contextlib import nullcontext
from datetime import date
import json
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from operator_app import engine, storage
from operator_app.glopro import GloProConnector
from test_engine import isolated_engine, source_file
from test_weekly import weekly_file


def output_folder(operators, path):
    conf = operators / "glopro.md"
    conf.write_text(conf.read_text().replace("rules: {}", f"rules: {{}}\nobsidian_output: '{path}'"))
    path.mkdir(exist_ok=True)
    return path


def test_import_selects_weekly_period_from_excel_and_updates_same_note_on_retry(isolated_engine, tmp_path):
    data, operators = isolated_engine
    output = output_folder(operators, tmp_path / "obsidian")
    target = output / "Активация — 22.09.2026.md"
    preserved = '22.09.2026\n\n### ООО "Другая"\nСумма проверена вручную.  \n\n'
    target.write_text(preserved + '### Китай\nСтарый раздел\n')
    for _ in range(2):
        source = weekly_file(tmp_path / "unrelated-filename.xlsx")
        nk = weekly_file(tmp_path / "nk.xlsx", 'ООО "НК АРТЭЛЬ"')
        run = engine.submit("glopro", "2026-09-22", "import", [
            {"path": str(source), "name": "unrelated-filename.xlsx"}, {"path": str(nk), "name": "nk.xlsx"}])
        assert run["status"] == "completed"
        text = target.read_text()
        assert preserved.split("\n\n", 1)[1] in text
        assert text.index("### Китай") < text.index('### ООО "Другая"')
        assert text.count("### Китай") == text.count('### ООО "НК АРТЭЛЬ"') == 1
        assert len(list(output.glob("*.md"))) == 1
        audit = json.loads((data / "runs" / run["id"] / "Проверка расчётов.json").read_text())
        assert all(r["period"] == ["2026-09-15", "2026-09-21"] for r in audit["reports"])
    book = load_workbook(data / "runs" / run["id"] / "Активация — 22.09.2026.xlsx")
    sheet = book.active
    assert sheet["G2"].value == 122.216 and sheet["G2"].number_format == "#,##0.000"
    assert sheet["E2"].value is None and sheet["E3"].value is None
    assert sheet["C3"].value == 1.01 and sheet["C3"].number_format == "#,##0.00"


@pytest.mark.parametrize("reason", ["old_period", "wrong_total"])
def test_mismatched_weekly_source_is_marked_review_and_does_not_publish(isolated_engine, tmp_path, reason):
    data, operators = isolated_engine
    output = output_folder(operators, tmp_path / "obsidian")
    target = output / "Активация — 22.09.2026.md"
    original = '22.09.2026\n\n### Другая фирма\nПроверено.\n'
    target.write_text(original)
    source = weekly_file(tmp_path / "bad.xlsx", start="18.09.2026" if reason == "old_period" else "15.09.2026")
    if reason == "wrong_total":
        book = load_workbook(source)
        for row in book.active:
            if row[1].value == "Итого по карте:":
                row[12].value = 99999
        book.save(source)
    run = engine.submit("glopro", "2026-09-22", "import", [{"path": str(source), "name": "bad.xlsx"}])
    assert run["status"] == "needs_review"
    assert target.read_text() == original
    assert run["report"] == ""
    assert list((data / "runs" / run["id"]).glob("На проверку*.zip"))


def test_friday_import_skips_both_weekly_firms_before_period_validation(isolated_engine, tmp_path):
    _, operators = isolated_engine
    output = output_folder(operators, tmp_path / "obsidian")
    inputs = [weekly_file(tmp_path / "china.xlsx"), weekly_file(tmp_path / "nk.xlsx", 'ООО "НК АРТЭЛЬ"'),
              source_file(tmp_path / "other.xlsx")]
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(p), "name": p.name} for p in inputs])
    assert run["status"] == "completed"
    assert len(run["schedule_skipped_clients"]) == 2
    assert "Китай" not in run["report"] and "НК АРТЭЛЬ" not in run["report"]
    assert "100,000 л | 8 300,00 рублей" in run["report"]


def test_weekly_zero_replaces_stale_section_without_claiming_activation(isolated_engine, tmp_path):
    _, operators = isolated_engine
    output = output_folder(operators, tmp_path / "obsidian")
    target = output / "Активация — 22.09.2026.md"
    target.write_text('22.09.2026\n\n### Китай\nустаревшая сумма\n')
    source = weekly_file(tmp_path / "zero.xlsx", cards=[])
    run = engine.submit("glopro", "2026-09-22", "import", [{"path": str(source), "name": source.name}])
    assert run["status"] == "no_data"
    assert "устаревшая" not in target.read_text()
    assert "клиент нам 0,00" in target.read_text()


def test_weekly_duplicate_import_never_doubles_amount(isolated_engine, tmp_path):
    data, _ = isolated_engine
    a = weekly_file(tmp_path / "a.xlsx")
    b = weekly_file(tmp_path / "b.xlsx")
    run = engine.submit("glopro", "2026-09-22", "import", [{"path": str(p), "name": p.name} for p in (a, b)])
    assert run["status"] == "failed"
    assert "Повтор фирмы" in run["error"]
    assert not list((data / "runs" / run["id"]).glob("Активация*"))


@pytest.mark.parametrize("run_date,start,end,expected", [
    ("2026-09-22", "2026-09-18", "2026-09-21", {"78749": date(2026, 9, 15), "78756": date(2026, 9, 15), "1": date(2026, 9, 18)}),
    ("2026-09-25", "2026-09-22", "2026-09-24", {"1": date(2026, 9, 22)}),
    ("2026-09-01", "2026-08-28", "2026-08-31", {"78749": date(2026, 8, 25), "78756": date(2026, 8, 25), "1": date(2026, 8, 28)}),
])
def test_connector_uses_company_period_and_skips_friday_before_contracts(tmp_path, run_date, start, end, expected):
    connector = GloProConnector("fixture", "fixture")
    clients = [{"id": "78749", "name": "Китай"}, {"id": "78756", "name": 'ООО "НК АРТЭЛЬ"'}, {"id": "1", "name": 'ООО "ТЕСТ"'}]
    events = []
    with patch.object(connector, "_browser_page", return_value=nullcontext(object())), \
         patch.object(connector, "_login"), \
         patch.object(connector, "_list_clients", return_value=clients), \
         patch.object(connector, "_contracts", return_value=[{"id": "10"}]) as contracts, \
         patch.object(connector, "_prepare_report", return_value={"checked": True}) as prepare, \
         patch.object(connector, "_download", return_value={}):
        sources = connector.download_reports(start, end, tmp_path, run_date=run_date, progress=events.append)
    assert {call.args[1]["id"] for call in contracts.call_args_list} == set(expected)
    assert {call.args[1]["id"]: call.args[3] for call in prepare.call_args_list} == expected
    assert all(call.args[4] == date.fromisoformat(end) for call in prepare.call_args_list)
    assert {s["client_id"]: s["start"] for s in sources} == {k: v.isoformat() for k, v in expected.items()}
    assert len([e for e in events if e.get("reason") == "tuesday_only"]) == (2 if run_date == "2026-09-25" else 0)


def test_remote_weekly_contracts_merge_by_holder(isolated_engine, monkeypatch):
    _, _ = isolated_engine
    def download(conf, record, directory, progress):
        return [{"path": str(weekly_file(directory / f"{c}.xlsx", 'ООО "НК АРТЭЛЬ"')), "client": 'ООО "НК АРТЭЛЬ"',
                 "client_id": "78756", "contract_id": c} for c in ("a", "b")]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-22")
    assert run["status"] == "completed"
    assert run["report"].count("• Имя из Excel") == 1
    assert "2,01 л | 222,20 руб" in run["report"]

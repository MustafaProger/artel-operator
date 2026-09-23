from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo
import hashlib
import json
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from operator_app import config, engine, storage


class ImmediatePool:
    def submit(self, fn, *args):
        fn(*args)


class DeferredPool:
    def __init__(self):
        self.calls = []

    def submit(self, fn, *args):
        self.calls.append((fn, args))


def source_file(path, client='ООО "ТЕСТ"', tx_date=date(2026, 9, 17), *, empty=False, litres=100, supplier=8000, customer=8300):
    book = Workbook()
    ws = book.active
    ws.title = "transactions"
    ws.append([client])
    ws.append(["Дата", "Услуга", "Количество", "Стоимость ТО", "Стоимость", "Скидка, %"])
    if not empty:
        ws.append([tx_date, "ДТ ЭКТО", litres, supplier, customer, -1])
    book.save(path)
    return path


@pytest.fixture
def isolated_engine(tmp_path, monkeypatch):
    data = tmp_path / "data"
    operators = tmp_path / "operators"
    operators.mkdir()
    for module in (config, storage, engine):
        monkeypatch.setattr(module, "DATA", data)
    monkeypatch.setattr(config, "OPERATORS", operators)
    monkeypatch.setattr(engine, "POOL", ImmediatePool())
    monkeypatch.setattr(storage, "credentials", lambda: None)
    def reject_network(*args, **kwargs):
        raise AssertionError("A test attempted an external download")
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": reject_network})
    markdown = """---
id: glopro
name: GloPro
kind: glopro
enabled: true
supplier: Новое имя
schedule:
  days: [tue, fri]
  time: '09:00'
  timezone: Europe/Moscow
rules: {}
---
Проверяемая инструкция.
"""
    (operators / "glopro.md").write_text(markdown, encoding="utf-8")
    storage.init()
    return data, operators


def test_good_import_produces_named_audited_artifacts_and_valid_zip(isolated_engine, tmp_path):
    data, _ = isolated_engine
    source = source_file(tmp_path / "source.xlsx")
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(source), "name": 'ООО ТЕСТ — 15.09.2026–17.09.2026.xlsx'}])
    saved = storage.get_run(run["id"])
    assert saved["status"] == "completed"
    assert saved["period_start"] == "2026-09-15"
    assert saved["period_end"] == "2026-09-17"
    root = data / "runs" / run["id"]
    assert (root / "Активация — 18.09.2026.md").read_text() == saved["report"]
    assert "100,000 л | 8 300,00 рублей" in saved["report"]
    audit = json.loads((root / "Проверка расчётов.json").read_text())
    assert audit["sources"][0]["sha256"] == original_hash
    assert audit["reports"][0]["totals"]["supplier_basis"] == "8000.00"
    sheet = load_workbook(root / "Активация — 18.09.2026.xlsx", data_only=True).active
    assert sheet.cell(2, 4).value == 8300
    archive = next(root.glob("Готовый комплект*.zip"))
    with zipfile.ZipFile(archive) as bundle:
        assert bundle.testzip() is None
        assert "Активация — 18.09.2026.md" in bundle.namelist()
        assert "Проверка расчётов.json" in bundle.namelist()
        assert any(name.startswith("Файлы по фирмам/") for name in bundle.namelist())
        assert all(not name.startswith("/") and ".." not in Path(name).parts for name in bundle.namelist())
    for entry in saved["files"]:
        assert entry["sha256"] == hashlib.sha256((root / entry["name"]).read_bytes()).hexdigest()
    assert not source.exists(), "temporary upload must be removed"


def test_malformed_input_never_releases_partial_ready_report(isolated_engine, tmp_path):
    data, _ = isolated_engine
    valid = source_file(tmp_path / "valid.xlsx")
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not an excel workbook")
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(valid), "name": "valid.xlsx"}, {"path": str(bad), "name": "bad.xlsx"}])
    root = data / "runs" / run["id"]
    assert run["status"] == "needs_review"
    assert len(run["failures"]) == 1
    assert run["report"] == ""
    assert not list(root.glob("Активация*"))
    assert not list(root.glob("Готовый комплект*"))
    assert list(root.glob("На проверку*.zip"))
    assert (root / "Ошибки проверки.json").is_file()


def test_excluded_remote_client_is_skipped_before_calculation(isolated_engine, monkeypatch):
    data, operators = isolated_engine
    config_path = operators / "glopro.md"
    config_path.write_text(config_path.read_text().replace("rules: {}", 'excluded_clients: [{id: "2", name: "Китай"}]\nrules: {}'))

    def download(conf, record, directory, progress):
        good = source_file(directory / "included.xlsx")
        excluded = directory / "excluded.xlsx"
        excluded.write_bytes(b"not a calculable workbook")
        return [{"path": str(good), "client": 'ООО "ТЕСТ"', "client_id": "1"},
                {"path": str(excluded), "client": "Новое имя исключённой фирмы", "client_id": "2"}]

    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "completed"
    assert run["client_count"] == 1
    assert run["excluded_client_count"] == 1
    assert run["failures"] == []
    assert "Новое имя исключённой фирмы" not in run["report"]
    audit = json.loads((data / "runs" / run["id"] / "Проверка расчётов.json").read_text())
    assert audit["excluded_clients"][0]["client_id"] == "2"
    assert len(audit["calculated_sources"]) == 1


def test_same_name_different_id_is_not_excluded_and_still_blocks_bad_report(isolated_engine, monkeypatch):
    _, operators = isolated_engine
    config_path = operators / "glopro.md"
    config_path.write_text(config_path.read_text().replace("rules: {}", 'excluded_clients: [{id: "2", name: "Китай"}]\nrules: {}'))

    def download(conf, record, directory, progress):
        good = source_file(directory / "included.xlsx")
        bad = directory / "different-client.xlsx"
        bad.write_bytes(b"invalid workbook")
        return [{"path": str(good), "client_id": "1"},
                {"path": str(bad), "client": "Китай", "client_id": "3"}]

    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "needs_review"
    assert run["excluded_client_count"] == 0
    assert len(run["failures"]) == 1


def test_download_exclusions_are_preserved_in_report_audit(isolated_engine, monkeypatch):
    data, _ = isolated_engine

    def download(conf, record, directory, progress):
        for _ in range(2):
            progress({"stage": "skipped", "reason": "excluded_client", "client": "Китай", "client_id": "2"})
        return [{"path": str(source_file(directory / "included.xlsx")), "client_id": "1"}]

    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "completed"
    assert run["source_count"] == 1
    assert run["excluded_client_count"] == 1
    audit = json.loads((data / "runs" / run["id"] / "Проверка расчётов.json").read_text())
    assert audit["excluded_clients"] == run["excluded_clients"]


@pytest.mark.parametrize("fully_reversed", [False, True])
def test_no_net_transactions_do_not_claim_completed_activation(isolated_engine, tmp_path, fully_reversed):
    data, operators = isolated_engine
    output = tmp_path / "obsidian"
    conf = operators / "glopro.md"
    conf.write_text(conf.read_text().replace("rules: {}", f"rules: {{}}\nobsidian_output: '{output}'"))
    path = source_file(tmp_path / "empty.xlsx", empty=not fully_reversed)
    if fully_reversed:
        book = load_workbook(path)
        book.active.append([date(2026, 9, 17), "ДТ", -100, -8000, -8300, -1])
        book.save(path)
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(path), "name": "empty.xlsx"}])
    root = data / "runs" / run["id"]
    assert run["status"] != "completed"
    assert not list(root.glob("Активация*"))
    assert not list(root.glob("Готовый комплект*"))
    assert not list(output.glob("*.md"))


def test_existing_obsidian_report_is_preserved(isolated_engine, tmp_path):
    _, operators = isolated_engine
    output = tmp_path / "obsidian"
    output.mkdir()
    human_report = output / "Активация — 18.09.2026.md"
    human_report.write_text("Проверенный вручную отчёт", encoding="utf-8")
    conf = operators / "glopro.md"
    conf.write_text(conf.read_text().replace("rules: {}", f"rules: {{}}\nobsidian_output: '{output}'"))
    source = source_file(tmp_path / "input.xlsx")
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(source), "name": "input.xlsx"}])
    assert "Проверенный вручную отчёт" in human_report.read_text()
    assert '### ООО "ТЕСТ"' in human_report.read_text()
    assert len(list(output.glob("*.md"))) == 1


def test_schedule_submits_once_per_date_across_ticks_and_restart(isolated_engine, monkeypatch):
    monkeypatch.setattr(storage, "credentials", lambda: ("fake", "not-a-real-secret"))
    def download(conf, record, directory, progress):
        path = source_file(directory / "client.xlsx", tx_date=date.fromisoformat(record["period_start"]))
        return [{"path": str(path)}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    storage.set_setting("enabled_since:glopro", "2026-09-14T08:00:00+03:00")
    now = datetime(2026, 9, 18, 9, 10, tzinfo=ZoneInfo("Europe/Moscow"))
    engine.tick(now)
    engine.tick(now)
    storage.init()
    engine.tick(now)
    runs = storage.runs()
    assert len(runs) == 2
    assert {r["run_date"] for r in runs} == {"2026-09-15", "2026-09-18"}
    assert all(r["status"] == "completed" and r["trigger"] == "schedule" for r in runs)
    assert {r["run_date"]: (r["period_start"], r["period_end"]) for r in runs} == {"2026-09-15": ("2026-09-11", "2026-09-14"), "2026-09-18": ("2026-09-15", "2026-09-17")}


def test_interrupted_scheduled_run_preserves_date_and_does_not_duplicate(isolated_engine, monkeypatch):
    monkeypatch.setattr(storage, "credentials", lambda: ("fake", "not-a-real-secret"))
    pool = DeferredPool()
    monkeypatch.setattr(engine, "POOL", pool)
    storage.set_setting("enabled_since:glopro", "2026-09-18T08:00:00+03:00")
    now = datetime(2026, 9, 18, 9, 10, tzinfo=ZoneInfo("Europe/Moscow"))
    engine.tick(now)
    assert len(pool.calls) == 1
    assert storage.runs()[0]["status"] == "queued"
    storage.init()
    engine.tick(now)
    assert len(pool.calls) == 1
    record = storage.runs()[0]
    assert record["status"] == "failed"
    assert record["run_date"] == "2026-09-18"
    assert record["period_start"] == "2026-09-15"
    assert record["period_end"] == "2026-09-17"


def test_submit_snapshots_markdown_before_background_execution(isolated_engine, monkeypatch):
    _, operators = isolated_engine
    pool = DeferredPool()
    monkeypatch.setattr(engine, "POOL", pool)
    run = engine.submit("glopro", "2026-09-18")
    original_snapshot = pool.calls[0][1][0]["markdown"]
    path = operators / "glopro.md"
    path.write_text(path.read_text().replace("Новое имя", "Другой поставщик"))
    assert pool.calls[0][1][0]["markdown"] == original_snapshot
    assert pool.calls[0][1][0]["supplier"] == "Новое имя"
    assert storage.get_run(run["id"])["status"] == "queued"


def test_same_company_with_different_quotes_cannot_be_imported_twice(isolated_engine, tmp_path):
    data, _ = isolated_engine
    first = source_file(tmp_path / "first.xlsx", client='ООО "ТЕСТ"')
    second = source_file(tmp_path / "second.xlsx", client="ООО «ТЕСТ»")
    run = engine.submit("glopro", "2026-09-18", "import", [{"path": str(first), "name": "first.xlsx"}, {"path": str(second), "name": "second.xlsx"}])
    assert run["status"] == "failed"
    assert "Повтор фирмы" in run["error"]
    root = data / "runs" / run["id"]
    assert not list(root.glob("Активация*"))
    assert not list(root.glob("Готовый комплект*"))


def test_distinct_remote_contracts_merge_to_one_company_with_weighted_price(isolated_engine, monkeypatch):
    data, _ = isolated_engine
    def download(conf, record, directory, progress):
        first = source_file(directory / "contract-a.xlsx", litres=1, supplier=100, customer=110)
        second = source_file(directory / "contract-b.xlsx", client="ООО «ТЕСТ»", litres=99, supplier=8000, customer=8200)
        return [{"path": str(first), "client": 'ООО "ТЕСТ"', "client_id": "same-company", "contract_id": "a"}, {"path": str(second), "client": "ООО «ТЕСТ»", "client_id": "same-company", "contract_id": "b"}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "completed"
    assert run["client_count"] == 1
    assert run["source_count"] == 2
    assert "80,19 рублей" in run["report"]
    assert "100,000 л | 8 310,00 рублей" in run["report"]
    audit = json.loads((data / "runs" / run["id"] / "Проверка расчётов.json").read_text())
    assert len(audit["reports"]) == 1
    assert audit["reports"][0]["audit"]["contract_ids"] == ["a", "b"]
    assert len(audit["calculated_sources"]) == 2


@pytest.mark.parametrize("contract_id", ["a", None])
def test_repeated_or_unidentified_remote_contract_is_not_merged(isolated_engine, monkeypatch, contract_id):
    data, _ = isolated_engine
    def download(conf, record, directory, progress):
        first = source_file(directory / "first.xlsx")
        second = source_file(directory / "second.xlsx")
        return [{"path": str(first), "client_id": "same-company", "contract_id": "a"}, {"path": str(second), "client_id": "same-company", "contract_id": contract_id}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "failed"
    assert "Повтор фирмы" in run["error"]
    assert not list((data / "runs" / run["id"]).glob("Активация*"))


def test_distinct_legal_entities_with_same_short_name_are_not_merged(isolated_engine, monkeypatch):
    data, _ = isolated_engine
    def download(conf, record, directory, progress):
        first = source_file(directory / "ooo.xlsx", client='ООО "ТЕСТ"', litres=1, supplier=100, customer=110)
        second = source_file(directory / "ao.xlsx", client='АО «ТЕСТ»', litres=99, supplier=8000, customer=8200)
        return [{"path": str(first), "client": 'ООО "ТЕСТ"', "client_id": "ooo", "contract_id": "a"}, {"path": str(second), "client": "АО «ТЕСТ»", "client_id": "ao", "contract_id": "b"}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "completed"
    assert run["client_count"] == 2
    audit = json.loads((data / "runs" / run["id"] / "Проверка расчётов.json").read_text())
    assert [(r["client_id"], r["totals"]["litres"]) for r in audit["reports"]] == [("ao", "99.000"), ("ooo", "1.000")]
    assert {r["client_id"] for r in audit["reports"]} == {"ooo", "ao"}
    assert all(not r["audit"].get("merged") for r in audit["reports"])


def test_same_display_name_but_different_remote_client_ids_is_not_merged(isolated_engine, monkeypatch):
    data, _ = isolated_engine
    def download(conf, record, directory, progress):
        first = source_file(directory / "first.xlsx")
        second = source_file(directory / "second.xlsx")
        return [{"path": str(first), "client_id": "one", "contract_id": "a"}, {"path": str(second), "client_id": "two", "contract_id": "b"}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "failed"
    assert "Повтор фирмы" in run["error"]
    assert not list((data / "runs" / run["id"]).glob("Активация*"))


def test_contracts_without_proven_client_id_are_not_merged(isolated_engine, monkeypatch):
    def download(conf, record, directory, progress):
        first = source_file(directory / "first.xlsx")
        second = source_file(directory / "second.xlsx")
        return [{"path": str(first), "contract_id": "a"}, {"path": str(second), "contract_id": "b"}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "failed"
    assert "Повтор фирмы" in run["error"]


def test_other_operator_uses_own_processor_without_fuel_calculator(isolated_engine, monkeypatch):
    data, operators = isolated_engine
    markdown = (operators / "glopro.md").read_text().replace("id: glopro", "id: documents").replace("kind: glopro", "kind: documents").replace("enabled: true", "enabled: false")
    (operators / "documents.md").write_text(markdown)
    def download(conf, record, directory, progress):
        path = directory / "documents.json"
        path.write_text('{"documents": 2}')
        return [{"path": str(path)}]
    def process(conf, record, sources, root, progress, *, imported):
        assert imported is False
        assert json.loads(Path(sources[0]["path"]).read_text())["documents"] == 2
        assert record["run_date"] == "2026-09-18"
        (root / "Сверка документов.txt").write_text("Проверено два документа")
        return engine.ProcessResult(status="completed", report="Проверено два документа", metrics={"document_count": 2}, audit={"documents": 2})
    def reject_fuel(*args, **kwargs):
        pytest.fail("Another operator must never invoke the fuel calculator")
    monkeypatch.setitem(engine.HANDLERS, "documents", download)
    monkeypatch.setitem(engine.PROCESSORS, "documents", process)
    monkeypatch.setattr(engine, "calculate_workbook", reject_fuel)
    run = engine.submit("documents", "2026-09-18")
    assert run["status"] == "completed"
    assert run["document_count"] == 2
    assert run["report"] == "Проверено два документа"
    root = data / "runs" / run["id"]
    assert not list(root.glob("Активация*"))
    assert (root / "Исходные файлы" / "documents.json").is_file()
    audit = json.loads((root / "Проверка расчётов.json").read_text())
    assert audit["documents"] == 2
    assert audit["run_id"] == run["id"]
    with zipfile.ZipFile(next(root.glob("Готовый комплект*.zip"))) as archive:
        assert archive.testzip() is None
        assert "Сверка документов.txt" in archive.namelist()


def test_operator_with_downloader_but_no_processor_is_rejected_before_enqueue(isolated_engine, monkeypatch):
    monkeypatch.setattr(engine, "PROCESSORS", {})
    with pytest.raises(ValueError, match="получение и обработка"):
        engine.submit("glopro", "2026-09-18")
    assert storage.runs() == []


@pytest.mark.parametrize("invalid_result", [engine.ProcessResult(status="running"), engine.ProcessResult(status="completed", metrics={"id": "overwrite"})])
def test_processor_cannot_publish_invalid_status_or_overwrite_run_identity(isolated_engine, monkeypatch, invalid_result):
    def download(conf, record, directory, progress):
        path = directory / "source.txt"
        path.write_text("source")
        return [{"path": str(path)}]
    monkeypatch.setattr(engine, "HANDLERS", {"glopro": download})
    monkeypatch.setitem(engine.PROCESSORS, "glopro", lambda *args, **kwargs: invalid_result)
    run = engine.submit("glopro", "2026-09-18")
    assert run["status"] == "failed"
    assert run["id"] != "overwrite"
    assert storage.get_run(run["id"])["status"] == "failed"

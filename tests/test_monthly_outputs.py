"""Monthly totals reconcile to displayed firms and retain source precision."""
from copy import deepcopy
from decimal import Decimal
import json

import pytest
from openpyxl import load_workbook

from operator_app.calculator import CalculationError, calculate_workbook, merge_reports
from operator_app.monthly import write_monthly_outputs
from operator_app.weekly import _report
from test_calculator import tx, workbook


RECORD = {"period_start": "2026-09-01", "period_end": "2026-09-30",
          "plan": {"kind": "monthly_summary", "month": "2026-09"}}


def calculated(tmp_path, client="Фирма", *, litres="10", supplier="800", customer="830"):
    return calculate_workbook(workbook(tmp_path, [tx(client=client, litres=litres, supplier=supplier, customer=customer)]))


def output(tmp_path, reports, **kwargs):
    return write_monthly_outputs(reports, RECORD, tmp_path / "output", **kwargs)


def excel(result):
    return load_workbook(next(path for path in result["files"] if path.endswith(".xlsx")), data_only=False)


def test_all_firms_returns_special_and_verified_zero_are_present(tmp_path):
    ordinary = calculate_workbook(workbook(tmp_path, [tx(client="Янтарь"), tx(client="Янтарь", litres=-20, supplier=-1600, customer=-1660)]))
    special = calculated(tmp_path, "ПЕРФЕКТ", litres="20", supplier="1567.2", customer="1900")
    reversed_report = calculate_workbook(workbook(tmp_path, [tx(client="Возврат"), tx(client="Возврат", litres=-100, supplier=-8000, customer=-8300)]))
    result = output(tmp_path, [ordinary, special, reversed_report], empty_clients=[{"client": "Альфа", "client_id": "1", "reason": "no_operations"}],
                    clients=[{"name": "Альфа", "id": "1"}, {"name": "Янтарь"}, {"name": "ПЕРФЕКТ"}, {"name": "Возврат"}])
    rows = result["metrics"]["companies"]
    assert [row["client"] for row in rows] == ["Альфа", "Возврат", "ПЕРФЕКТ", "Янтарь"]
    assert [row["status"] for row in rows] == ["no_data", "fully_reversed", "ready", "ready"]
    assert result["metrics"]["client_count"] == 4
    assert result["metrics"]["active_client_count"] == 2
    assert result["metrics"]["no_data_client_count"] == 1
    assert result["metrics"]["fully_reversed_client_count"] == 1
    assert result["metrics"]["total_litres"] == "100.000"
    assert result["metrics"]["total_customer_amount"] == "8238.54"
    assert "01.09.2026–30.09.2026" in result["report"]
    assert "Платежи и начальный остаток не вычитаются" in result["report"]
    assert "Операции полностью возвращены" in result["report"]
    assert len(result["files"]) == 3
    json.dumps(result, ensure_ascii=False)
    book = excel(result)
    assert book.sheetnames == ["Сводка по фирмам", "Детализация", "Период и правила"]
    sheet = book.worksheets[0]
    assert sheet["B6"].value == 100
    assert sheet["C6"].value == 8238.54
    assert sheet["B2"].data_type == "n" and sheet["C2"].data_type == "n"
    book.close()


def test_china_nk_use_raw_litres_and_customer_p_not_supplier_cost(tmp_path):
    reports = []
    for client, kind, holder in [("Китай", "china", ""), ('ООО "НК АРТЭЛЬ"', "nk_artel", "Держатель")]:
        reports.append(_report(client, kind, "Поставщик", [
            {"fuel": "ДТ ЭКТО", "holder": holder, "litres": Decimal("1.0054"), "supplier_basis": Decimal("900"), "customer_total": Decimal("80.015")},
            {"fuel": "АИ-95", "holder": holder, "litres": Decimal("2.0044"), "supplier_basis": Decimal("700"), "customer_total": Decimal("40.015")},
        ], {"operation_count": 2}))
    assert reports[1]["holders"][0]["litres"] == "1.01"
    result = output(tmp_path, reports)
    assert result["metrics"]["total_litres"] == "6.020"
    assert result["metrics"]["total_customer_amount"] == "240.06"
    nk = result["audit"]["monthly_summary"]["companies"][1]
    actual = next(item for item in nk["details"] if item["fuel"] == "ДТ ЭКТО")
    assert actual["litres"] == "1.005"
    assert actual["holder"] == "Держатель"
    assert nk["raw_totals"]["litres"] == "3.0098"
    assert sum(Decimal(row["customer_total"]) for row in nk["details"]) == Decimal(nk["customer_total"])
    assert any(row["kind"] == "rounding" for row in nk["details"])


def test_grand_total_adds_company_rounded_amounts_and_audits_difference(tmp_path):
    reports = [calculated(tmp_path, client=name, litres="1.0004", supplier="0.005", customer="0.005") for name in ("А", "Б")]
    result = output(tmp_path, reports)
    assert result["metrics"]["total_customer_amount"] == "0.02"
    assert result["metrics"]["total_litres"] == "2.000"
    audit = result["audit"]["monthly_summary"]
    assert audit["raw_totals"] == {"litres": "2.0008", "customer_total": "0.010"}
    assert Decimal(audit["rounding_adjustment"]["customer_total"]) == Decimal("0.010")
    book = excel(result)
    sheet = book.worksheets[0]
    assert Decimal(str(sheet["C2"].value)) + Decimal(str(sheet["C3"].value)) == Decimal(str(sheet["C4"].value))
    book.close()


def test_detail_rounding_adjustments_reconcile_each_firm(tmp_path):
    report = calculate_workbook(workbook(tmp_path, [tx(client="А", litres="1.0004", supplier="0.005", customer="0.005"),
                                                tx(client="А", fuel="АИ-95", litres="1.0004", supplier="0.005", customer="0.005")]))
    result = output(tmp_path, [report])
    company = result["audit"]["monthly_summary"]["companies"][0]
    rounding = next(row for row in company["details"] if row["kind"] == "rounding")
    assert rounding["litres"] == "0.001"
    assert rounding["customer_total"] == "-0.01"
    for key in ("litres", "customer_total"):
        assert sum(Decimal(row[key]) for row in company["details"]) == Decimal(company[key])
    assert "Корректировка округления" in result["report"]


def test_configurable_special_markup_applied_only_once_after_merging(tmp_path):
    rules = {"special_clients": ["Особая"], "special_customer_multiplier": "1.035"}
    reports = [calculate_workbook(workbook(tmp_path, [tx(client="Особая", litres=1, supplier="0.25", customer=9)]), rules) for _ in range(2)]
    merged = merge_reports(reports, rules)
    result = output(tmp_path, [merged], rules=rules)
    assert result["metrics"]["total_customer_amount"] == "0.52"
    assert result["audit"]["monthly_summary"]["raw_totals"]["customer_total"] == "0.51750"
    assert "Особая" in result["report"] and "× 1,035" in result["report"]
    assert "ПЕРФЕКТ" not in result["report"]


def test_negative_return_only_month_keeps_negative_amounts(tmp_path):
    report = calculated(tmp_path, litres="-10", supplier="-800", customer="-830")
    result = output(tmp_path, [report])
    assert result["metrics"]["total_litres"] == "-10.000"
    assert result["metrics"]["total_customer_amount"] == "-830.00"


@pytest.mark.parametrize("name", ['=HYPERLINK("https://example.com","click")', "+cmd", "-danger", "@SUM(A1:A9)", "<b>Фирма</b>|[link](url)"])
def test_untrusted_company_holder_and_fuel_are_xlsx_text(tmp_path, name):
    report = _report(name, "nk_artel", "Supplier", [{"holder": name, "fuel": name, "litres": Decimal(1),
                                                     "supplier_basis": Decimal(2), "customer_total": Decimal(3)}], {"operation_count": 1})
    result = output(tmp_path, [report])
    book = excel(result)
    for sheet, cells in [(book.worksheets[0], ["A2"]), (book.worksheets[1], ["A2", "B2", "C2"])]:
        for ref in cells:
            assert sheet[ref].value == name
            assert sheet[ref].data_type == "s"
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()
    if "<b>" in name:
        assert "<b>" not in result["report"]
        assert "\\|" in result["report"] and "\\[link\\]" in result["report"]


def test_verified_empty_month_still_outputs_complete_zero_summary(tmp_path):
    result = output(tmp_path, [], clients=[{"name": "А", "id": "1"}, {"name": "Б", "id": "2"}],
                    empty_clients=[{"client": "Б", "client_id": "2", "reason": "no_operations"}, {"client": "А", "client_id": "1", "reason": "no_operations"}])
    assert result["metrics"]["client_count"] == 2
    assert result["metrics"]["total_litres"] == "0.000"
    assert result["metrics"]["total_customer_amount"] == "0.00"


def test_empty_contract_does_not_zero_or_duplicate_company_with_operations(tmp_path):
    report = calculated(tmp_path, "А")
    report["client_id"] = "1"
    result = output(tmp_path, [report], empty_clients=[{"client": "А", "client_id": "1", "reason": "no_operations"}])
    assert result["metrics"]["client_count"] == 1
    assert result["metrics"]["total_customer_amount"] == "830.00"


@pytest.mark.parametrize("empty", [[], [{"client": "А", "client_id": "2", "reason": "no_operations"}]])
def test_manifest_firm_without_verified_coverage_fails_before_writing(tmp_path, empty):
    with pytest.raises(CalculationError, match="нет проверенного"):
        output(tmp_path, [], clients=[{"name": "А", "id": "1"}], empty_clients=empty)
    assert not (tmp_path / "output").exists()


@pytest.mark.parametrize("first,last", [("2026-09-02", "2026-09-30"), ("2026-09-01", "2026-10-01"), ("2026-02-01", "2026-02-27")])
def test_partial_or_cross_month_period_refused(tmp_path, first, last):
    with pytest.raises(CalculationError, match="полный календарный"):
        write_monthly_outputs([], {"period_start": first, "period_end": last}, tmp_path)


def test_leap_february_is_a_full_month(tmp_path):
    result = write_monthly_outputs([], {"period_start": "2028-02-01", "period_end": "2028-02-29"}, tmp_path)
    assert "01.02.2028–29.02.2028" in result["report"]


def test_mismatched_period_failed_status_and_duplicate_reports_refused(tmp_path):
    report = calculated(tmp_path)
    mismatched = deepcopy(report)
    mismatched["period"] = ["2026-09-15", "2026-09-21"]
    with pytest.raises(CalculationError, match="период фирмы"):
        output(tmp_path, [mismatched])
    failed = {**report, "status": "needs_review"}
    with pytest.raises(CalculationError, match="не прошёл"):
        output(tmp_path, [failed])
    with pytest.raises(CalculationError, match="Повтор фирмы"):
        output(tmp_path, [report, report])


def test_nonzero_zero_status_and_broken_detail_refused(tmp_path):
    report = calculated(tmp_path)
    with pytest.raises(CalculationError, match="нулевой статус"):
        output(tmp_path, [{**report, "status": "fully_reversed"}])
    report["audit"]["raw_fuels"][0]["customer_total"] = "1"
    with pytest.raises(CalculationError, match="детализация не сходится"):
        output(tmp_path, [report])

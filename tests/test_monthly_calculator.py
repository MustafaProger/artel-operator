from decimal import Decimal

import pytest
from openpyxl import load_workbook

from operator_app.calculator import CalculationError, calculate_workbook, merge_reports
from test_weekly import weekly_file


MONTH = "2026-09"


def monthly_source(tmp_path, client='ООО "ТЕСТ"', *, name="monthly.xlsx", cards=None, **kwargs):
    return weekly_file(tmp_path / name, client, cards=cards, start="01.09.2026", end="30.09.2026", **kwargs)


def calculate(path, **settings):
    # Wednesday is deliberately outside both Tuesday and month-close runs.
    return calculate_workbook(path, {"report_month": MONTH, "run_date": "2026-10-07", **settings})


def test_monthly_ordinary_uses_client_amount_and_all_cards_without_supplier_prices(tmp_path):
    path = monthly_source(tmp_path, cards=[
        ("virtual", "Первый", [("ДТ", "10.0004", "800", "793.04")]),
        ("plastic", "Второй", [("Дт", "20.0004", "1600", "1568"), ("АИ-95", "2", "120", "119")]),
        ("return", "Второй", [("ДТ ЭКТО", "-1", "-80", "-78.40")]),
    ])
    book = load_workbook(path)
    for row in book.active:
        if row[7].value in ("Дебет", "Возврат"):
            row[18].value = -1 if row[9].value == "10.0004" else -2
    book.save(path)
    report = calculate(path)
    assert report["report_month"] == MONTH
    assert report["totals"] == {"litres": "31.001", "supplier_basis": "2440.00", "customer_total": "2401.64"}
    assert report["audit"]["raw_totals"]["litres"] == "31.0008"
    assert report["card_type"] is None
    assert all("supplier_price" not in fuel for fuel in report["fuels"])
    assert report["audit"]["verified_period"] == ["2026-09-01T00:00:00+03:00", "2026-09-30T23:59:59+03:00"]
    assert {check["scope"].split()[0] for check in report["audit"]["reconciliation"]} == {"услуги", "карты", "файла"}


@pytest.mark.parametrize("client", ['ООО "ПЕРФЕКТ"', 'ООО "АЛЬФА-ЗАПАД"'])
def test_monthly_special_uses_configured_markup_once_from_supplier_basis(tmp_path, client):
    report = calculate(monthly_source(tmp_path, client), special_customer_multiplier="1.025")
    assert report["totals"]["customer_total"] == "126.54"
    assert report["audit"]["raw_totals"]["customer_total"] == "126.53625"
    assert report["audit"]["source_raw_totals"]["customer_total"] == "111.10"
    assert report["audit"]["special_client"] is True


@pytest.mark.parametrize("client", ['ООО "ПЕРФЕКТ"', 'ООО "АЛЬФА-ЗАПАД"'])
def test_monthly_special_gasoline_still_requires_defined_rule(tmp_path, client):
    path = monthly_source(tmp_path, client, cards=[("1", "Имя", [("АИ-95", "1", "100", "99")])])
    with pytest.raises(CalculationError, match="исключение для бензина"):
        calculate(path)


@pytest.mark.parametrize("client,kind", [("Китай", "china"), ('ООО "НК АРТЭЛЬ"', "nk_artel")])
def test_monthly_weekly_firms_use_full_month_and_keep_source_payable(tmp_path, client, kind):
    report = calculate(monthly_source(tmp_path, client))
    assert report["weekly_kind"] == kind
    assert report["totals"]["customer_total"] == "111.10"
    assert report["audit"]["raw_totals"]["customer_total"] == "111.10"
    assert report["report_month"] == MONTH


@pytest.mark.parametrize("client", ['ООО "ТЕСТ"', "Китай", 'ООО "НК АРТЭЛЬ"'])
@pytest.mark.parametrize("change", ["missing_period", "short_period", "partial_day", "missing_file", "missing_card", "missing_service", "wrong_file", "wrong_card", "wrong_service"])
def test_all_monthly_firms_require_full_period_and_reconciled_totals(tmp_path, client, change):
    path = monthly_source(tmp_path, client)
    book = load_workbook(path)
    ws = book.active
    scope = change.split("_")[-1]
    expected_label = {"file": "Итого по отчету:", "card": "Итого по карте:", "service": "Итого по ДТ:"}.get(scope)
    for row in ws:
        if row[1].value == "Период:":
            if change == "missing_period":
                for cell in row:
                    cell.value = None
            elif change == "short_period":
                row[4].value = "с 01.09.2026 00:00:00 по 29.09.2026 23:59:59"
            elif change == "partial_day":
                row[4].value = "с 01.09.2026 01:00:00 по 30.09.2026 23:59:59"
        if expected_label and row[1].value == expected_label:
            if change.startswith("missing_"):
                for cell in row:
                    cell.value = None
            else:
                row[15].value = "111.104"
    book.save(path)
    with pytest.raises(CalculationError):
        calculate(path)


@pytest.mark.parametrize("settings", [
    {"report_month": "2026-9"}, {"report_month": "2026-13"}, {"report_month": "0000-01"},
    {"date_from": "2026-09-02"}, {"date_to": "2026-09-29"},
])
def test_monthly_invalid_month_or_partial_boundaries_are_rejected(tmp_path, settings):
    with pytest.raises(CalculationError):
        calculate(monthly_source(tmp_path), **settings)


def test_monthly_leap_february_is_a_full_calendar_month(tmp_path):
    path = weekly_file(tmp_path / "leap.xlsx", 'ООО "ТЕСТ"', start="01.02.2024", end="29.02.2024")
    report = calculate(path, report_month="2024-02")
    assert report["audit"]["verified_period"][1] == "2024-02-29T23:59:59+03:00"


def test_monthly_ordinary_does_not_need_supplier_card_discount_column(tmp_path):
    path = monthly_source(tmp_path)
    book = load_workbook(path)
    book.active["S2"] = None
    book.save(path)
    assert calculate(path)["totals"]["customer_total"] == "111.10"


@pytest.mark.parametrize("client", ['ООО "ТЕСТ"', 'ООО "ПЕРФЕКТ"', "Китай", 'ООО "НК АРТЭЛЬ"'])
def test_monthly_merges_contracts_from_raw_values_before_rounding(tmp_path, client):
    reports = []
    for index in range(2):
        path = monthly_source(tmp_path, client, name=f"contract{index}.xlsx", cards=[
            (str(index), "Один держатель", [("ДТ", "1.0004", "0.005", "0.0049")]),
        ])
        report = calculate(path)
        report["contract_id"] = str(index)
        reports.append(report)
    merged = merge_reports(reports, {"special_customer_multiplier": "1.02"})
    assert merged["totals"]["litres"] == "2.001"
    assert merged["totals"]["customer_total"] == "0.01"
    assert Decimal(merged["audit"]["raw_totals"]["customer_total"]) == Decimal("0.0102" if "ПЕРФЕКТ" in client else "0.0098")
    assert merged["report_month"] == MONTH
    assert merged["audit"]["contract_ids"] == ["0", "1"]


def test_monthly_merging_refuses_missing_month_marker_and_other_period(tmp_path):
    first = calculate(monthly_source(tmp_path))
    second = calculate(monthly_source(tmp_path, name="second.xlsx"))
    second.pop("report_month")
    with pytest.raises(CalculationError, match="разных периодов"):
        merge_reports([first, second])
    second["report_month"] = "2026-08"
    with pytest.raises(CalculationError, match="разных периодов"):
        merge_reports([first, second])


def test_monthly_empty_and_fully_reversed_sources_are_distinct(tmp_path):
    empty = calculate(monthly_source(tmp_path, name="empty.xlsx", cards=[]))
    assert empty["status"] == "no_data"
    reversed_report = calculate(monthly_source(tmp_path, name="returns.xlsx", cards=[
        ("1", "Имя", [("ДТ", "1", "80", "79"), ("ДТ", "-1", "-80", "-79")]),
    ]))
    assert reversed_report["status"] == "fully_reversed"
    assert reversed_report["totals"]["customer_total"] == "0.00"

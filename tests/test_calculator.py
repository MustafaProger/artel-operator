from datetime import date

import pytest
from openpyxl import Workbook

from operator_app.calculator import CalculationError, calculate_workbook, company_identity, merge_reports, render_report


HEADERS = ["Клиент", "Дата", "Услуга", "Количество", "Стоимость ТО", "Стоимость", "Скидка, %"]


def workbook(tmp_path, rows, headers=None, metadata=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "transactions"
    for row in metadata or []:
        sheet.append(row)
    sheet.append(headers or HEADERS)
    for row in rows:
        sheet.append(row)
    path = tmp_path / "transactions.xlsx"
    book.save(path)
    return path


def tx(fuel="ДТ ЭКТО", litres=100, supplier=8000, customer=8300, discount=-1, client='ООО "ТЕСТ"'):
    return [client, date(2026, 9, 17), fuel, litres, supplier, customer, discount]


def test_supplier_and_customer_columns_are_distinct_and_refunds_are_signed(tmp_path):
    path = workbook(tmp_path, [tx(), tx("ДТ", 50, 4500, 4700), tx("ДТ ЭКТО", -20, -1600, -1660)])
    report = calculate_workbook(path)
    assert report["status"] == "ready"
    assert report["card_type"] == "virtual"
    assert report["fuels"] == [{"fuel": "ДТ ЭКТО", "litres": "130.000", "supplier_basis": "10900.00", "customer_total": "11340.00", "supplier_price": "83.01", "rows": [2, 3, 4]}]
    assert "130,000 л | 11 340,00 рублей" in render_report([report], "18.09.2026")
    assert "10 900" not in render_report([report], "18.09.2026")


@pytest.mark.parametrize("client", ['ООО "ПЕРФЕКТ"', 'ООО «АЛЬФА-ЗАПАД»'])
def test_special_clients_use_supplier_basis_for_customer_markup(tmp_path, client):
    path = workbook(tmp_path, [tx(litres=20, supplier=1567.2, customer=1900, discount=-2, client=client)])
    report = calculate_workbook(path)
    assert report["card_type"] == "virtual"
    assert report["fuels"][0]["customer_total"] == "1598.54"
    assert report["fuels"][0]["supplier_price"] == "77.58"


def test_mixed_types_refused_instead_of_averaged(tmp_path):
    path = workbook(tmp_path, [tx(discount=-1), tx(discount=-2)])
    with pytest.raises(CalculationError, match="смешаны"):
        calculate_workbook(path)


def test_headers_are_named_not_fixed_position_and_summaries_not_counted(tmp_path):
    headers = ["Услуга", "Стоимость", "Дата операции", "Скидка, %", "Количество", "Клиент", "Стоимость ТО"]
    rows = [["ДТ", 8300, date(2026, 9, 17), -2, 100, 'ООО "ТЕСТ"', 8000], ["Итого", 8300, None, None, 100, None, 8000], headers]
    report = calculate_workbook(workbook(tmp_path, rows, headers))
    assert report["fuels"][0]["supplier_price"] == "80.40"
    assert report["totals"]["customer_total"] == "8300.00"
    assert report["audit"]["skipped_summary_rows"] == [3]
    assert report["audit"]["skipped_header_rows"] == [4]


def test_duplicate_header_is_ambiguous(tmp_path):
    path = workbook(tmp_path, [], HEADERS + ["Стоимость"])
    with pytest.raises(CalculationError, match="Неоднозначный"):
        calculate_workbook(path, {"client": "Тест"})


def test_money_discount_is_not_percentage(tmp_path):
    path = workbook(tmp_path, [tx(discount=-300)], HEADERS[:-1] + ["Скидка"])
    with pytest.raises(CalculationError, match="колонками"):
        calculate_workbook(path)


def test_percentage_number_format(tmp_path):
    path = workbook(tmp_path, [tx(discount=-0.01)])
    from openpyxl import load_workbook
    book = load_workbook(path)
    book.active["G2"].number_format = "0.00%"
    book.save(path)
    assert calculate_workbook(path)["card_type"] == "virtual"


def test_fully_reversed_operations_not_divided_by_zero(tmp_path):
    path = workbook(tmp_path, [tx(), tx(litres=-100, supplier=-8000, customer=-8300)])
    report = calculate_workbook(path)
    assert report["status"] == "fully_reversed"
    assert report["fuels"] == []
    assert report["totals"] == {"litres": "0.000", "supplier_basis": "0.00", "customer_total": "0.00"}
    assert "Операции полностью возвращены." in render_report([report], "2026-09-18")


def test_nonzero_amount_with_zero_net_litres_refused(tmp_path):
    path = workbook(tmp_path, [tx(), tx(litres=-100, supplier=-7900, customer=-8300)])
    with pytest.raises(CalculationError, match="литраж равен нулю"):
        calculate_workbook(path)


def test_empty_table_is_no_data(tmp_path):
    report = calculate_workbook(workbook(tmp_path, []), {"client": "Тест"})
    assert report["status"] == "no_data"
    assert report["audit"]["transaction_rows"] == []


def test_gasoline_only_does_not_guess_card_type(tmp_path):
    report = calculate_workbook(workbook(tmp_path, [tx("АИ-92 ЭКТО", 10, 600, 620, -2)]))
    assert report["card_type"] is None
    assert report["fuels"][0]["supplier_price"] == "60.60"
    assert '- АИ-92 ЭКТО — 10,000 л | 620,00 руб | 60,60 руб/л' in render_report([report], "2026-09-18")


def test_special_client_gasoline_is_ambiguous(tmp_path):
    path = workbook(tmp_path, [tx("АИ-92 ЭКТО", client="ПЕРФЕКТ")])
    with pytest.raises(CalculationError, match="исключение для бензина"):
        calculate_workbook(path)


@pytest.mark.parametrize("fuel", ["Мойка", "Газ", "Неизвестно"])
def test_unknown_services_require_explicit_rules(tmp_path, fuel):
    with pytest.raises(CalculationError, match="неизвестная услуга"):
        calculate_workbook(workbook(tmp_path, [tx(fuel)]))


def test_client_mismatch_refused(tmp_path):
    with pytest.raises(CalculationError, match="другая фирма"):
        calculate_workbook(workbook(tmp_path, [tx()]), {"client": "Другая"})


def test_period_mismatch_refused_not_silently_filtered(tmp_path):
    with pytest.raises(CalculationError, match="вне запрошенного периода"):
        calculate_workbook(workbook(tmp_path, [tx()]), {"date_from": "2026-09-18", "date_to": "2026-09-21"})


def test_configurable_rules_and_russian_numbers(tmp_path):
    report = calculate_workbook(workbook(tmp_path, [tx(litres="100,000", supplier="8 000,00", customer="8\xa0300,00", discount="-1,00%")]), {"virtual_diesel_multiplier": "0.98"})
    assert report["fuels"][0]["supplier_price"] == "78.40"


def test_exact_concise_markdown(tmp_path):
    report = calculate_workbook(workbook(tmp_path, [tx()]))
    assert render_report([report], "2026-09-18") == '18.09.2026\n\n### ООО "ТЕСТ"\nПоставщик: "Новое имя"  \n79,20 рублей  \n100,000 л | 8 300,00 рублей\n'


def test_equal_transactions_are_not_deduplicated(tmp_path):
    report = calculate_workbook(workbook(tmp_path, [tx(), tx()]))
    assert report["totals"]["litres"] == "200.000"


def test_unknown_discount_not_guessed(tmp_path):
    with pytest.raises(CalculationError, match="не определён"):
        calculate_workbook(workbook(tmp_path, [tx(discount=-5)]))


def test_refund_sign_mismatch_refused(tmp_path):
    with pytest.raises(CalculationError, match="знаки"):
        calculate_workbook(workbook(tmp_path, [tx(litres=-100)]))


def test_real_layout_metadata_cards_and_footer_are_not_transactions(tmp_path):
    headers = [None] * 20
    for index, name in {3: "Дата", 7: "Операция", 8: "Услуга", 9: "Количество", 12: "Стоимость ТО", 15: "Стоимость", 16: "Скидка", 18: "Скидка, %"}.items():
        headers[index] = name
    title = [None, 'ООО "ТЕСТ"']
    card_row = [None, "Карта: пример", None, "Держатель: пример"]
    transaction = [None] * 20
    for index, value in {3: date(2026, 9, 17), 7: "Дебет", 8: "ДТ ЭКТО", 9: 100, 12: 8000, 15: 8300, 16: -300, 18: -1}.items():
        transaction[index] = value
    total = [None, "Итого по отчету:"] + [None] * 18
    total[9], total[12], total[15] = 100, 8000, 8300
    rows = [card_row, transaction, [None, "Итого по карте:"], total, [None, "Общий итог по услугам"], [None, "Услуга", "Количество", "Стоимость ТО", "Стоимость"], [None, "ДТ ЭКТО", 100, 8000, 8300], [None, date(2026, 9, 19)]]
    result = calculate_workbook(workbook(tmp_path, rows, headers, [title]))
    assert result["client"] == 'ООО "ТЕСТ"'
    assert result["totals"]["customer_total"] == "8300.00"
    assert result["audit"]["columns"]["supplier_basis"] == 13
    assert result["audit"]["skipped_metadata_rows"] == [3]


def test_wrong_xlsx_dimension_does_not_hide_transactions(tmp_path):
    import zipfile
    import re
    path = workbook(tmp_path, [tx()])
    with zipfile.ZipFile(path) as z:
        files = {name: z.read(name) for name in z.namelist()}
    files["xl/worksheets/sheet1.xml"] = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', files["xl/worksheets/sheet1.xml"])
    with zipfile.ZipFile(path, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    assert calculate_workbook(path)["totals"]["litres"] == "100.000"


def test_transaction_comment_with_total_word_is_still_counted(tmp_path):
    row = tx() + ["Итого оплачено"]
    result = calculate_workbook(workbook(tmp_path, [row], HEADERS + ["Комментарий"]))
    assert result["totals"]["litres"] == "100.000"


def test_merge_contracts_recomputes_weighted_price_from_sums(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx(litres=1, supplier=100, customer=110)]))
    second = calculate_workbook(workbook(tmp_path, [tx(litres=99, supplier=8000, customer=8200)]))
    first["contract_id"], second["contract_id"] = "a", "b"
    merged = merge_reports([first, second])
    assert merged["fuels"][0]["supplier_price"] == "80.19"
    assert merged["fuels"][0]["customer_total"] == "8310.00"
    assert merged["fuels"][0]["litres"] == "100.000"
    assert merged["audit"]["contract_ids"] == ["a", "b"]
    assert len(merged["audit"]["source_reports"]) == 2


def test_merge_contracts_preserves_unrounded_values(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx(litres="1.0004", supplier="1.004", customer="1.006")]))
    second = calculate_workbook(workbook(tmp_path, [tx(litres="1.0004", supplier="1.004", customer="1.006")]))
    merged = merge_reports([first, second])
    assert first["fuels"][0]["litres"] == "1.000"
    assert merged["fuels"][0]["litres"] == "2.001"
    assert merged["fuels"][0]["supplier_basis"] == "2.01"
    assert merged["fuels"][0]["customer_total"] == "2.01"
    assert merged["audit"]["raw_totals"]["customer_total"] == "2.012"


def test_merge_special_client_applies_customer_markup_once_after_sum(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx(litres=1, supplier="0.25", customer=5, client="ПЕРФЕКТ")]))
    second = calculate_workbook(workbook(tmp_path, [tx(litres=1, supplier="0.25", customer=5, client="ПЕРФЕКТ")]))
    assert first["fuels"][0]["customer_total"] == "0.26"
    merged = merge_reports([first, second])
    assert merged["fuels"][0]["customer_total"] == "0.51"
    assert merged["fuels"][0]["supplier_price"] == "0.25"


def test_merge_no_data_contracts_remains_no_data(tmp_path):
    first = calculate_workbook(workbook(tmp_path, []), {"client": "ТЕСТ", "card_type": "virtual"})
    second = calculate_workbook(workbook(tmp_path, []), {"client": "ТЕСТ", "card_type": "plastic"})
    merged = merge_reports([first, second])
    assert merged["status"] == "no_data"
    assert merged["fuels"] == []


def test_merge_ignores_no_data_card_type_but_preserves_ready_peer(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx()]))
    second = calculate_workbook(workbook(tmp_path, []), {"client": 'ООО «ТЕСТ»', "card_type": "plastic"})
    merged = merge_reports([first, second])
    assert merged["status"] == "ready"
    assert merged["card_type"] == "virtual"
    assert merged["totals"]["litres"] == "100.000"


def test_merge_mixed_card_types_rejected(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx(discount=-1)]))
    second = calculate_workbook(workbook(tmp_path, [tx(discount=-2)]))
    with pytest.raises(CalculationError, match="смешаны"):
        merge_reports([first, second])


@pytest.mark.parametrize("change", ["client", "supplier"])
def test_merge_inconsistent_client_or_supplier_rejected(tmp_path, change):
    first = calculate_workbook(workbook(tmp_path, [tx()]))
    second = calculate_workbook(workbook(tmp_path, [tx()]))
    second[change] = "Другой"
    with pytest.raises(CalculationError, match="разные фирмы"):
        merge_reports([first, second])


def test_merge_full_refund_across_contracts(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx()]))
    second = calculate_workbook(workbook(tmp_path, [tx(litres=-100, supplier=-8000, customer=-8300)]))
    result = merge_reports([first, second])
    assert result["status"] == "fully_reversed"
    assert result["fuels"] == []


def test_merge_duplicate_fuel_mapping_rejected(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx()]))
    first["audit"]["raw_fuels"].append({**first["audit"]["raw_fuels"][0], "fuel": "ДТ"})
    with pytest.raises(CalculationError, match="Повтор топлива"):
        merge_reports([first])


def test_company_identity_preserves_legal_form_and_normalizes_quote_style():
    assert company_identity('  ооо   « ТЕСТ » ') == company_identity('ООО "ТЕСТ"')
    assert company_identity('ООО "ТЕСТ"') != company_identity('АО "ТЕСТ"')


def test_workbook_cannot_be_attributed_to_different_legal_entity(tmp_path):
    path = workbook(tmp_path, [tx(client='АО "ТЕСТ"')])
    with pytest.raises(CalculationError, match="другая фирма"):
        calculate_workbook(path, {"client": 'ООО "ТЕСТ"'})


def test_merge_rejects_different_legal_entities_with_same_short_name(tmp_path):
    first = calculate_workbook(workbook(tmp_path, [tx(client='ООО "ТЕСТ"')]))
    second = calculate_workbook(workbook(tmp_path, [tx(client='АО "ТЕСТ"')]))
    with pytest.raises(CalculationError, match="разные фирмы"):
        merge_reports([first, second])

"""Discovery, pagination, source reconciliation and identity regression tests."""
from copy import deepcopy
from datetime import date
from pathlib import Path
from unittest.mock import Mock
import json

import pytest
from openpyxl import load_workbook

from operator_app import config, engine
from operator_app.yandex import download_employee
from operator_app.yandex_cabinet import collect_orders, discover_employees, CabinetError
from operator_app.yandex_reports import read_report, phone_hash, write_outputs
from operator_app.yandex_publication import merge_yandex_sections
from test_engine import isolated_engine
from test_yandex import START, END, STAFF, source, modify, yandex_config

COMPANY = 'c' * 32


def order(n=1, uid='a'*32, name='Джибраил', phone='71111111111', **changes):
    result = dict(id=f'{n:032x}', user_id=uid, client_id=COMPANY,
                  created_at=f'2026-09-15T12:00:{60-n:02d}+03:00', status='Completed',
                  user_info=dict(fullname=name, phone=phone), liters_filled='30',
                  final_price='3141.90', currency='RUB', is_active=False, is_deleted=True)
    result.update(changes)
    return result


class Pages:
    company_id = COMPANY
    def __init__(self, pages):
        self.pages = iter(pages)
        self.cursors = []
    def orders_page(self, start, end, cursor, limit):
        assert (start, end) == (START, END)
        self.cursors.append(cursor)
        return next(self.pages)


def page(rows, cursor='', limit=30, **extra):
    return dict(orders=rows, cursor=cursor, limit=limit, sorting_order='desc', **extra)


def manifest(rows):
    return collect_orders(Pages([page(rows)]), START, END)


def workbook(path, raw):
    source(path)
    for cell, value in [('B8', raw['user_info']['fullname']), ('C8',raw['user_info']['phone']),
                        ('D8',raw['id']),('G8',float(raw['liters_filled'])),
                        ('J8',float(raw['final_price'])),('D5',float(raw['final_price'])),('J9',float(raw['final_price']))]:
        modify(path,cell,value)
    return path


def test_all_pages_discover_third_and_deleted_person_without_roster():
    rows = [order(), order(2,STAFF[1]['user_id'],'Эльдар','72222222222'),
            order(3,'d'*32,'Аслан','73333333333')]
    api = Pages([page(rows[:2],'next',2),page(rows[2:],'',2)])
    full = collect_orders(api,START,END,limit=2)
    people = discover_employees(full,STAFF)
    assert api.cursors == ['', 'next']
    assert [p['name'] for p in people] == ['Аслан','Джибриль','Эльдар']
    assert full['order_count'] == 3 and full['complete']
    assert people[0]['user_id']=='d'*32
    assert full['pages']==[dict(number=1,count=2,has_next=True),dict(number=2,count=1,has_next=False)]


@pytest.mark.parametrize('pages',[
    [page([order()], 'x', 1),page([order()], '', 1)],  # repeated order
    [page([order()], 'x', 1),page([order(2)], 'x', 1)], # looping cursor
    [page([order()], '', 1)], # truncation without terminal cursor
    [page([order()], 'x', 2),page([], 'y', 2)],
    [dict(orders=[],limit=2,sorting_order='desc')], # missing cursor isn't empty success
    [page([], '', 2, total=1)],
    [page([order(2)], 'x', 2),page([order()], '', 2)], # backwards pagination
    [page([order()], 'x', 2,total=2),page([order(2)],'',2,total=3)],
])
def test_incomplete_pagination_is_never_empty_success(pages):
    with pytest.raises(CabinetError):collect_orders(Pages(pages),START,END,limit=pages[0]['limit'])


@pytest.mark.parametrize('changes',[
    dict(user_id=''),dict(client_id='f'*32),dict(created_at='2026-09-15T12:00:00Z'),
    dict(created_at='2026-09-22T00:00:00+03:00'),dict(currency='USD'),
    dict(status='Refunded',liters_filled='0',final_price='0'),
    dict(status='Unknown',liters_filled='0',final_price='0'),
    dict(liters_filled='-30',final_price='-3141.90'),
    dict(user_info=dict(fullname='Иван',phone='')),dict(final_price=None),
])
def test_unknown_status_identity_or_finance_blocks_result(changes):
    with pytest.raises(CabinetError):manifest([order(**changes)])


def test_same_names_different_ids_have_distinct_files_and_captions(tmp_path):
    raw=[order(1,'a'*32,'Иван','71111111111'),order(2,'b'*32,'Иван','72222222222')]
    full=manifest(raw);people=discover_employees(full)
    assert len({p['label'] for p in people})==2
    reports=[read_report(workbook(tmp_path/f'{i}.xlsx',o),START,END,'ООО НК АРТЭЛЬ',expected_orders=p['orders']) for i,(o,p) in enumerate(zip(raw,people))]
    text=write_outputs(reports,tmp_path,date(2026,9,22))
    assert text.count('на склад (Иван)')==2
    assert f'### Иван [{"a"*32}]' in text and f'### Иван [{"b"*32}]' in text
    sheet=load_workbook(tmp_path/'Яндекс — 22.09.2026.xlsx').active
    assert [sheet.cell(i,6).value for i in (2,3)]==['a'*32,'b'*32]


def test_same_name_and_phone_still_distinct_by_cabinet_id(tmp_path):
    raw=[order(1,'a'*32,'Иван'),order(2,'b'*32,'Иван')]
    people=discover_employees(manifest(raw))
    assert len(people)==2
    for i,(o,p) in enumerate(zip(raw,people)):
        report=read_report(workbook(tmp_path/f'{i}.xlsx',o),START,END,'ООО НК АРТЭЛЬ',expected_orders=p['orders'])
        assert report['user_id']==p['user_id']


def test_renamed_employee_keeps_id_and_latest_source_name(tmp_path):
    raw=[order(1,'a'*32,'Новое имя'),order(2,'a'*32,'Старое имя')]
    people=discover_employees(manifest(raw))
    assert len(people)==1 and people[0]['source_name']=='Новое имя'
    path=workbook(tmp_path/'renamed.xlsx',raw[0]);b=load_workbook(path);ws=b.active
    ws.insert_rows(9);ws.append([])
    ws.cell(9,1,'15.09.2026');ws.cell(9,2,'Старое имя');ws.cell(9,3,'71111111111');ws.cell(9,4,raw[1]['id']);ws.cell(9,5,'АИ-100');ws.cell(9,7,30);ws.cell(9,8,'Завершён');ws.cell(9,10,3141.9)
    ws['D4']=2;ws['D5']=6283.8;ws['J10']=6283.8;b.save(path)
    report=read_report(path,START,END,'ООО НК АРТЭЛЬ',expected_orders=people[0]['orders'])
    assert report['name']=='Новое имя' and report['order_count']==2


def test_dynamic_import_unknown_employee_and_identity_limit(tmp_path):
    o=order(3,'d'*32,'Аслан','73333333333');path=workbook(tmp_path/'third.xlsx',o)
    report=read_report(path,START,END,'ООО НК АРТЭЛЬ',aliases=STAFF)
    assert report['name']=='Аслан' and report['identity_basis']=='xlsx_phone_sha256'
    aliases=[dict(STAFF[0],phone_sha256=phone_hash('73333333333')),dict(STAFF[1],phone_sha256=phone_hash('73333333333'))]
    with pytest.raises(ValueError,match='несколькими ID'):read_report(path,START,END,'ООО НК АРТЭЛЬ',aliases=aliases)


@pytest.mark.parametrize('cell,value',[('D8','f'*32),('B8','Другой'),('C8','79999999999'),('G8',31),('J8',3000)])
def test_excel_must_match_each_cabinet_order(tmp_path,cell,value):
    raw=order();p=workbook(tmp_path/'r.xlsx',raw);modify(p,cell,value)
    with pytest.raises(ValueError):read_report(p,START,END,'ООО НК АРТЭЛЬ',expected_orders=manifest([raw])['orders'])


def test_skip_cancelled_only_employee_before_any_report_request(tmp_path):
    raw=order(status='Cancelled',liters_filled='0',final_price='0')
    employee=discover_employees(manifest([raw]))[0];api=Mock();browser=Mock();events=[]
    assert download_employee(api,browser,tmp_path,employee,START,END,'ООО НК АРТЭЛЬ',[],events.append) is None
    assert not api.mock_calls and not browser.mock_calls and not list(tmp_path.iterdir())
    assert events[0]['reason']=='no_operations'


def test_dynamic_full_engine_and_missing_third_blocks_publication(isolated_engine,tmp_path,monkeypatch):
    data,operators=isolated_engine;out=tmp_path/'notes';yandex_config(operators,out)
    raw=[order(),order(2,STAFF[1]['user_id'],'Эльдар','72222222222'),order(3,'d'*32,'Аслан','73333333333')]
    full=manifest(raw)
    def handler(conf,record,directory,progress):
        progress(dict(stage='discovered',manifest=full))
        progress(dict(activity_check_complete=True,order_count=3,employee_ids=[o['user_id'] for o in raw]))
        return [dict(path=str(workbook(directory/f'{i}.xlsx',o)),user_id=o['user_id'],report_id=f'{i+5:032x}') for i,o in enumerate(raw)]
    monkeypatch.setitem(engine.HANDLERS,'yandex',handler)
    run=engine.submit('yandex','2026-09-22');assert run['status']=='completed'
    assert run['report'].index('### Аслан')<run['report'].index('### Джибриль')<run['report'].index('### Эльдар')
    assert run['active_client_count']==3
    note=out/'Яндекс Заправки — 22.09.2026.md';before=note.read_text()
    def incomplete(*args):return handler(*args)[:-1]
    monkeypatch.setitem(engine.HANDLERS,'yandex',incomplete)
    run=engine.submit('yandex','2026-09-22')
    assert run['status']=='needs_review' and not run['report'] and note.read_text()==before
    assert not list((data/'runs'/run['id']).glob('Яндекс — *'))


def test_complete_empty_week_no_xlsx(isolated_engine,tmp_path,monkeypatch):
    data,operators=isolated_engine;yandex_config(operators)
    def handler(conf,record,directory,progress):
        progress(dict(stage='discovered',manifest=manifest([])))
        progress(dict(stage='skipped',reason='no_operations',client=conf['company']))
        progress(dict(activity_check_complete=True,order_count=0,employee_ids=[]))
        return []
    monkeypatch.setitem(engine.HANDLERS,'yandex',handler)
    run=engine.submit('yandex','2026-09-22')
    assert run['status']=='no_data'
    assert not list((data/'runs'/run['id']).rglob('*.xlsx'))


def test_note_updates_by_identity_preserve_other_sections_and_disambiguate_incremental_namesakes(tmp_path):
    incoming=[]
    for i,phone in enumerate(['71111111111','72222222222'],1):
        raw=order(i,uid=str(i)*32,name='Иван',phone=phone)
        report=read_report(workbook(tmp_path/f'{i}.xlsx',raw),START,END,'ООО НК АРТЭЛЬ')
        incoming.append(write_outputs([report],tmp_path,date(2026,9,22)))
    other='## Заметки\nРучной текст\n\n### Лукойл\nДругой раздел\n'
    combined=merge_yandex_sections(other,incoming[0])
    combined=merge_yandex_sections(combined,incoming[1])
    assert combined.count('# Яндекс Заправки — 22.09.2026')==1
    assert combined.count('на склад (Иван)')==2 and other.strip() in combined
    assert '### Иван [phone-' in combined
    assert merge_yandex_sections(combined,incoming[0])==combined
    changed=incoming[0].replace('Иван','Переименованный')
    updated=merge_yandex_sections(combined,changed)
    assert updated.count('на склад (Иван)')==1 and updated.count('на склад (Переименованный)')==1


def test_config_without_staff_list_and_namesakes(isolated_engine):
    _,operators=isolated_engine;md=yandex_config(operators)
    import yaml
    data=yaml.safe_load(md.split('---')[1]);data.pop('employees')
    config.parse_operator('---\n'+yaml.safe_dump(data,allow_unicode=True)+'---\nВсе сотрудники')
    data['employee_aliases']=[dict(e,name='Иван') for e in STAFF]
    config.parse_operator('---\n'+yaml.safe_dump(data,allow_unicode=True)+'---\nПодписи')


def test_full_last_page_requires_terminal_empty_page():
    api=Pages([page([order()], 'last', 1),page([], '', 1)])
    full=collect_orders(api,START,END,limit=1)
    assert full['order_count']==1 and api.cursors==['','last']


def test_row_litres_use_observed_xlsx_precision(tmp_path):
    raw=order(liters_filled='48.037')
    path=workbook(tmp_path/'precision.xlsx',raw);modify(path,'G8',48.04)
    expected=manifest([raw])['orders']
    report=read_report(path,START,END,'ООО НК АРТЭЛЬ',expected_orders=expected)
    assert report['litres']=='48.04'
    modify(path,'G8',48.03)
    with pytest.raises(ValueError,match='Строка Excel'):read_report(path,START,END,'ООО НК АРТЭЛЬ',expected_orders=expected)


def test_request_is_unfiltered_includes_last_fraction_and_generate_has_idempotency():
    from operator_app.yandex_cabinet import Cabinet,API,ORDERS,REPORTS
    context=Mock();response=Mock();response.status=200
    response.json.return_value=dict(id=COMPANY,name='ООО НК АРТЭЛЬ')
    response.request.all_headers.return_value={'x-csrf-token':'test-token','cookie':'private','x-yataxi-selected-corp-client-id':COMPANY}
    api=Cabinet(context,response,'ООО НК АРТЭЛЬ')
    result=context.request.fetch.return_value;result.status=200;result.json.return_value={}
    api.orders_page(START,END,'',30)
    args,kwargs=context.request.fetch.call_args
    from urllib.parse import urlparse,parse_qs
    assert urlparse(args[0]).path==ORDERS
    query=parse_qs(urlparse(args[0]).query)
    assert query['since_datetime']==['2026-09-15T00:00:00+03:00']
    assert query['till_datetime']==['2026-09-21T23:59:59.999999+03:00']
    assert json.loads(kwargs['data'])=={} and 'cookie' not in kwargs['headers']
    api.request(REPORTS+'generate',dict(service='tanker'))
    first=context.request.fetch.call_args.kwargs['headers']['x-idempotency-token']
    api.request(REPORTS+'generate',dict(service='tanker'))
    assert context.request.fetch.call_args.kwargs['headers']['x-idempotency-token']!=first
    result.status=500
    with pytest.raises(CabinetError,match='отсутствие операций не подтверждено'):api.orders_page(START,END,'',30)


def test_source_name_collision_sanitizing_does_not_overwrite(tmp_path):
    from operator_app.yandex_reports import label_reports
    reports=[dict(name='Иван/Петров',employee_id='a'*32),dict(name='Иван\\Петров',employee_id='b'*32)]
    label_reports(reports)
    assert len({r['label'] for r in reports})==2
    assert all('/' not in r['label'] and '\\' not in r['label'] for r in reports)


def test_ambiguous_legacy_note_does_not_publish_files(isolated_engine,tmp_path,monkeypatch):
    monkeypatch.setitem(engine.HANDLERS,"yandex",Mock(side_effect=AssertionError("network")))
    data,operators=isolated_engine;notes=tmp_path/'notes';notes.mkdir();yandex_config(operators,notes)
    note=notes/'Яндекс Заправки — 22.09.2026.md'
    old='# Яндекс Заправки — 22.09.2026\n\n### Джибриль\n15.09 - 21.09\nЯндекс заправки\nна склад (Джибриль) 10,00 л на сумму закупки 1,00 рублей\n'
    note.write_text(old)
    path=source(tmp_path/'dj.xlsx')
    run=engine.submit('yandex','2026-09-22','import',[dict(path=str(path),name=path.name)])
    assert run['status']=='needs_review' and not run['report'] and note.read_text()==old
    assert not list((data/'runs'/run['id']).glob('Яндекс — *'))


def test_snapshot_change_after_download_blocks_completion(tmp_path,monkeypatch):
    from unittest.mock import MagicMock
    from operator_app import yandex
    full=manifest([order()]);changed=manifest([order(),order(2,'d'*32,'Аслан','73333333333')])
    session=tmp_path/'session.json';session.write_text('{}')
    monkeypatch.setattr(yandex.yandex_connection,'session_path',lambda:session)
    manager=MagicMock();monkeypatch.setattr(yandex,'sync_playwright',lambda:manager)
    monkeypatch.setattr(yandex,'Cabinet',Mock())
    monkeypatch.setattr(yandex,'verify_page',Mock())
    monkeypatch.setattr(yandex,'collect_orders',Mock(side_effect=[full,changed]))
    downloader=Mock(return_value=dict(path='one.xlsx'));monkeypatch.setattr(yandex,'download_employee',downloader)
    events=[]
    with pytest.raises(yandex.YandexError,match='изменились'):
        yandex.download_reports(dict(company='ООО НК АРТЭЛЬ'),dict(run_date='2026-09-22'),tmp_path,events.append)
    assert downloader.call_count==1
    assert not any(e.get('activity_check_complete') for e in events)


def test_import_note_rebinding_requires_order_ids_not_just_phone(tmp_path):
    raw=order(name='Иван');path=workbook(tmp_path/'same.xlsx',raw)
    known=read_report(path,START,END,'ООО НК АРТЭЛЬ',expected_orders=manifest([raw])['orders'])
    offline=read_report(path,START,END,'ООО НК АРТЭЛЬ')
    before=write_outputs([known],tmp_path,date(2026,9,22))
    incoming=write_outputs([offline],tmp_path,date(2026,9,22))
    after=merge_yandex_sections(before,incoming)
    assert after.count('на склад (Иван)')==1 and 'yandex-employee:'+raw['user_id'] in after
    other=order(2,'b'*32,'Иван');workbook(path,other)
    report=read_report(path,START,END,'ООО НК АРТЭЛЬ')
    with pytest.raises(ValueError,match='только по телефону'):
        merge_yandex_sections(before,write_outputs([report],tmp_path,date(2026,9,22)))


def test_engine_namesakes_keep_two_originals(isolated_engine,tmp_path,monkeypatch):
    data,operators=isolated_engine;yandex_config(operators)
    monkeypatch.setitem(engine.HANDLERS,'yandex',Mock(side_effect=AssertionError('network')))
    files=[workbook(tmp_path/f'{i}.xlsx',order(i+1,str(i)*32,'Иван',str(i+3)*11)) for i in range(2)]
    run=engine.submit('yandex','2026-09-22','import',[dict(path=str(p),name=p.name) for p in files])
    assert run['status']=='completed' and run['active_client_count']==2
    originals=list((data/'runs'/run['id']/'Исходные файлы').glob('*.xlsx'))
    assert len(originals)==2 and len({p.name for p in originals})==2
    assert all('Яндекс. Иван [phone-' in p.name for p in originals)
